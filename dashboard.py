# NOTE: This file is reconstructed from the two dashboard code sections
# supplied in the conversation, with integrated COLOR / AUTHENTICITY support.
# Replace your current dashboard file with this file.

from pathlib import Path
from datetime import datetime
import json
import re

import pandas as pd
import streamlit as st


# ============================================================
# PROJECT PATHS
# ============================================================

BASE_DIR = Path(__file__).resolve().parent

DATA_DIR = BASE_DIR / "data"
OUTPUTS_DIR = BASE_DIR / "outputs"
DOCS_DIR = BASE_DIR / "docs"

DATASET_PATH = DATA_DIR / "raw" / "saffron.csv"

README_PATH = BASE_DIR / "README.md"
REFERENCE_DOCX_PATH = DOCS_DIR / "saffron-05.docx"
CONCLUSION_PATH = DOCS_DIR / "conclusion_fa.md"
RUN_OUTPUT_PATH = BASE_DIR / "run_output.txt"

ROBUST_SUMMARY_PATH = BASE_DIR / "robust_evaluation_summary.csv"
ROBUST_RESULTS_PATH = BASE_DIR / "robust_evaluation_results.csv"
NOVELTY_PATH = BASE_DIR / "novelty_detection_results.csv"
PREDICTIONS_PATH = BASE_DIR / "sample_predictions.csv"
DECISION_PATH = BASE_DIR / "decision_engine_results.csv"
SHAP_PATH = BASE_DIR / "shap_feature_importance.csv"
FINAL_REPORT_PATH = BASE_DIR / "final_report.csv"

CROSS_YEAR_PATH = OUTPUTS_DIR / "cross_year_origin_results.csv"
YEAR_PREDICTION_PATH = OUTPUTS_DIR / "harvest_year_prediction_cv.csv"
YEAR_EFFECT_PATH = OUTPUTS_DIR / "within_region_year_effect.csv"
YEAR_EFFECT_SUMMARY_PATH = OUTPUTS_DIR / "harvest_year_effect_summary.csv"

HISTORY_DIR = OUTPUTS_DIR / "dashboard_history"

REGION_YEAR_OUTPUT_DIR = OUTPUTS_DIR / "region_year_analysis"
REGION_DISTRIBUTION_PATH = REGION_YEAR_OUTPUT_DIR / "region_distribution.csv"
REGIONAL_MODEL_RESULTS_PATH = REGION_YEAR_OUTPUT_DIR / "regional_model_results.csv"
REGION_YEAR_SIMILARITY_PATH = REGION_YEAR_OUTPUT_DIR / "region_year_similarity.csv"
REGION_YEAR_PAIRWISE_PATH = REGION_YEAR_OUTPUT_DIR / "region_year_pairwise_distances.csv"
REGION_YEAR_SUMMARY_PATH = REGION_YEAR_OUTPUT_DIR / "analysis_summary.csv"

COLOR_DASHBOARD_JSON = (
    BASE_DIR / "reports" / "color_dashboard" / "color_dashboard_result.json"
)
COLOR_DASHBOARD_CSV = (
    BASE_DIR / "reports" / "color_dashboard" / "color_dashboard_samples.csv"
)
COLOR_EDA_ALL_PLOT = (
    BASE_DIR / "reports" / "color_eda" / "all_spectra.png"
)
COLOR_EDA_REGION_PLOT = (
    BASE_DIR / "reports" / "color_eda" / "color_region_5_9ppm.png"
)
COLOR_PCA_REGION_PLOT = (
    BASE_DIR / "reports" / "color_pca_color_region" / "pca_scores_color_region.png"
)

PCA_SCORES_PLOT = BASE_DIR / "pca_scores.png"
PCA_SCREE_PLOT = BASE_DIR / "pca_scree_plot.png"
NOVELTY_PLOT = BASE_DIR / "novelty_detection.png"
SHAP_PLOT = BASE_DIR / "shap_summary.png"
YEAR_PCA_PLOT = OUTPUTS_DIR / "year_pca_plot.png"
YEAR_REGION_PCA_PLOT = OUTPUTS_DIR / "year_region_pca_plot.png"


# ============================================================
# PAGE CONFIG
# ============================================================

st.set_page_config(
    page_title="Saffron NMR Research Dashboard",
    page_icon="🧪",
    layout="wide",
    initial_sidebar_state="expanded",
)


# ============================================================
# CSS
# ============================================================

st.markdown(
    """
    <style>
    html, body, .stApp, [data-testid="stAppViewContainer"], [data-testid="stMain"] {
        direction: rtl !important;
    }
    .main, .block-container {
        direction: rtl !important;
        text-align: right !important;
    }
    .block-container {
        max-width: 1500px;
        padding-top: 1.5rem;
        padding-bottom: 3rem;
    }
    body, p, span, div, label, button, input, textarea {
        font-family: Tahoma, Arial, sans-serif !important;
    }
    [data-testid="stMarkdownContainer"] {
        direction: rtl !important;
        text-align: right !important;
    }
    [data-testid="stMarkdownContainer"] p {
        direction: rtl !important;
        text-align: right !important;
        line-height: 2;
    }
    [data-testid="stMarkdownContainer"] h1,
    [data-testid="stMarkdownContainer"] h2,
    [data-testid="stMarkdownContainer"] h3,
    [data-testid="stMarkdownContainer"] h4,
    [data-testid="stMarkdownContainer"] h5,
    [data-testid="stMarkdownContainer"] h6,
    [data-testid="stMarkdownContainer"] ul,
    [data-testid="stMarkdownContainer"] ol,
    [data-testid="stMarkdownContainer"] li {
        direction: rtl !important;
        text-align: right !important;
    }
    [data-testid="stMetric"] {
        direction: rtl !important;
        text-align: center !important;
        background: white;
        border: 1px solid #e4eae6;
        border-radius: 16px;
        padding: 15px 10px;
        min-height: 115px;
        box-shadow: 0 4px 14px rgba(0, 0, 0, 0.04);
    }
    [data-testid="stMetricLabel"] {
        direction: rtl !important;
        text-align: center !important;
    }
    [data-testid="stMetricValue"] {
        direction: ltr !important;
        text-align: center !important;
        font-weight: 800 !important;
    }
    [data-testid="stTabs"] {
        direction: rtl !important;
    }
    [data-testid="stTabs"] [role="tablist"] {
        direction: rtl !important;
        justify-content: flex-start !important;
        gap: 6px;
        flex-wrap: wrap;
    }
    [data-testid="stTabs"] button {
        direction: rtl !important;
        text-align: right !important;
        font-family: Tahoma, Arial, sans-serif !important;
        font-weight: 600 !important;
        font-size: 13px !important;
    }
    [data-testid="stTabs"] button div,
    [data-testid="stTabs"] button p,
    [data-testid="stTabs"] button span {
        direction: rtl !important;
        text-align: right !important;
    }
    /* README = tab 12 */
    [data-testid="stTabs"] [data-baseweb="tab-panel"]:nth-of-type(12) {
        direction: ltr !important;
        text-align: left !important;
    }
    [data-testid="stTabs"] [data-baseweb="tab-panel"]:nth-of-type(12) [data-testid="stMarkdownContainer"] {
        direction: ltr !important;
        text-align: left !important;
    }
    [data-testid="stTabs"] [data-baseweb="tab-panel"]:nth-of-type(12) [data-testid="stMarkdownContainer"] p,
    [data-testid="stTabs"] [data-baseweb="tab-panel"]:nth-of-type(12) [data-testid="stMarkdownContainer"] h1,
    [data-testid="stTabs"] [data-baseweb="tab-panel"]:nth-of-type(12) [data-testid="stMarkdownContainer"] h2,
    [data-testid="stTabs"] [data-baseweb="tab-panel"]:nth-of-type(12) [data-testid="stMarkdownContainer"] h3,
    [data-testid="stTabs"] [data-baseweb="tab-panel"]:nth-of-type(12) [data-testid="stMarkdownContainer"] h4,
    [data-testid="stTabs"] [data-baseweb="tab-panel"]:nth-of-type(12) [data-testid="stMarkdownContainer"] li {
        direction: ltr !important;
        text-align: left !important;
    }
    [data-testid="stCodeBlock"] {
        direction: ltr !important;
        text-align: left !important;
    }
    [data-testid="stCodeBlock"] pre, [data-testid="stCodeBlock"] code {
        direction: ltr !important;
        text-align: left !important;
        font-family: Consolas, "Courier New", monospace !important;
        font-size: 13px !important;
        line-height: 1.6 !important;
    }
    [data-testid="stAlert"], [data-testid="stAlert"] * {
        direction: rtl !important;
        text-align: right !important;
    }
    [data-testid="stSidebar"], [data-testid="stSidebar"] * {
        direction: rtl !important;
        text-align: right !important;
    }
    [data-testid="stButton"] button {
        direction: rtl !important;
        text-align: center !important;
        font-family: Tahoma, Arial, sans-serif !important;
    }
    [data-testid="stSelectbox"] {
        direction: rtl !important;
        text-align: right !important;
    }
    [data-testid="stDataFrame"] { direction: ltr !important; }
    [data-testid="stImage"] { direction: ltr !important; }
    [data-testid="stExpander"] { direction: rtl !important; }
    </style>
    """,
    unsafe_allow_html=True,
)


# ============================================================
# FILE IO
# ============================================================

@st.cache_data(show_spinner=False)
def read_text(path_string: str) -> str:
    path = Path(path_string)
    if not path.exists():
        return ""
    for encoding in ("utf-8", "utf-8-sig", "cp1252"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
        except Exception:
            return ""
    return ""


@st.cache_data(show_spinner=False)
def read_csv(path_string: str) -> pd.DataFrame:
    path = Path(path_string)
    if not path.exists():
        return pd.DataFrame()
    try:
        return pd.read_csv(path)
    except Exception:
        return pd.DataFrame()


@st.cache_data(show_spinner=False)
def read_json(path_string: str):
    path = Path(path_string)
    if not path.exists():
        return {}
    try:
        with open(path, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return {}


# ============================================================
# HELPERS
# ============================================================

def fmt(value, decimals=4):
    try:
        return f"{float(value):.{decimals}f}"
    except Exception:
        return "—"


def fmt_pm(mean, std):
    if pd.isna(mean):
        return "—"
    if pd.isna(std):
        return fmt(mean)
    return f"{fmt(mean)} ± {fmt(std)}"


def normalize_bool(value):
    return str(value).strip().lower() in {"true", "1", "yes"}


def file_time(path: Path):
    if not path.exists():
        return "وجود ندارد"
    try:
        return datetime.fromtimestamp(path.stat().st_mtime).strftime("%Y-%m-%d %H:%M:%S")
    except Exception:
        return "نامشخص"


def find_column(df: pd.DataFrame, candidates):
    for candidate in candidates:
        if candidate in df.columns:
            return candidate
    return None


def show_image(path: Path, caption: str):
    if path.exists():
        st.image(str(path), caption=caption, width="stretch")
    else:
        st.info(f"نمودار موجود نیست: {path.name}")


# ============================================================
# DATASET INFORMATION
# ============================================================

@st.cache_data(show_spinner=False)
def get_dataset_info():
    result = {
        "initial_records": 0,
        "independent_samples": 0,
        "metadata_columns": 0,
        "spectral_points": 0,
        "groups": 0,
        "group_counts": pd.DataFrame(),
        "missing_values": 0,
        "invalid_values": 0,
        "duplicate_extra_rows": 0,
        "year_counts": pd.DataFrame(),
        "years": [],
        "region_year": pd.DataFrame(),
        "min_ppm": None,
        "max_ppm": None,
    }

    df = read_csv(str(DATASET_PATH))
    if df.empty:
        return result

    result["initial_records"] = len(df)

    metadata_names = {
        "name", "group", "region_code", "harvest_year", "sample_id", "id"
    }

    metadata_columns = [
        col for col in df.columns if col.lower() in metadata_names
    ]

    spectral_columns = [
        col for col in df.columns if col not in metadata_columns
    ]

    result["metadata_columns"] = len(metadata_columns)
    result["spectral_points"] = len(spectral_columns)

    if "group" in df.columns:
        group_counts = (
            df["group"]
            .astype(str)
            .value_counts()
            .sort_index()
            .rename_axis("Group")
            .reset_index(name="Samples")
        )
        result["group_counts"] = group_counts
        result["groups"] = len(group_counts)

    result["missing_values"] = int(df.isna().sum().sum())

    if spectral_columns:
        numeric_values = df[spectral_columns].apply(
            pd.to_numeric,
            errors="coerce",
        )
        result["invalid_values"] = int(
            numeric_values.isna().sum().sum()
        )

        ppm = pd.to_numeric(
            pd.Index(spectral_columns),
            errors="coerce",
        )
        ppm = ppm[~pd.isna(ppm)]

        if len(ppm) > 0:
            result["min_ppm"] = float(ppm.min())
            result["max_ppm"] = float(ppm.max())

        try:
            duplicate_mask = df[spectral_columns].duplicated(keep="first")
            result["duplicate_extra_rows"] = int(duplicate_mask.sum())
        except Exception:
            pass

    result["independent_samples"] = max(
        0,
        result["initial_records"] - result["duplicate_extra_rows"],
    )

    year_values = []

    if "harvest_year" in df.columns:
        year_values = pd.to_numeric(
            df["harvest_year"], errors="coerce"
        ).tolist()
    elif "HarvestYear" in df.columns:
        year_values = pd.to_numeric(
            df["HarvestYear"], errors="coerce"
        ).tolist()
    elif "name" in df.columns:
        for value in df["name"].astype(str).str.strip():
            match = re.match(r"^(\d{2})-", value)
            if not match:
                year_values.append(None)
                continue
            prefix = int(match.group(1))
            year_values.append(1400 + prefix if prefix <= 29 else 1300 + prefix)
    else:
        year_values = [None] * len(df)

    year_series = pd.Series(year_values)

    year_counts = (
        year_series.dropna()
        .astype(int)
        .value_counts()
        .sort_index()
        .rename_axis("HarvestYear")
        .reset_index(name="Samples")
    )

    result["year_counts"] = year_counts
    result["years"] = year_counts["HarvestYear"].astype(int).tolist()

    if "group" in df.columns and not year_series.empty:
        temp = pd.DataFrame({"Group": df["group"], "HarvestYear": year_series})
        result["region_year"] = pd.crosstab(
            temp["Group"], temp["HarvestYear"]
        ).sort_index()

    return result


# ============================================================
# PROJECT METRICS
# ============================================================

def get_metrics():
    dataset = get_dataset_info()

    result = {
        "samples": dataset["independent_samples"],
        "groups": dataset["groups"],
        "ood": 0,
        "top_k": "—",
        "accuracy": "—",
        "accuracy_std": "—",
        "ba": "—",
        "ba_std": "—",
        "f1": "—",
        "f1_std": "—",
        "chance_ba": "—",
        "status": "NOT RELIABLE",
    }

    novelty = read_csv(str(NOVELTY_PATH))
    if not novelty.empty:
        status_column = find_column(novelty, ["NoveltyStatus", "Status"])
        if status_column:
            try:
                result["ood"] = int(
                    novelty[status_column]
                    .astype(str)
                    .str.contains(r"NOVEL|OOD", case=False, regex=True)
                    .sum()
                )
            except Exception:
                pass

    robust = read_csv(str(ROBUST_SUMMARY_PATH))
    required = {
        "TopK", "AccuracyMean", "AccuracyStd", "BalancedAccuracyMean",
        "BalancedAccuracyStd", "F1MacroMean", "F1MacroStd",
        "ChanceBalancedAccuracy", "AboveChanceCandidate",
    }

    if not robust.empty and required.issubset(robust.columns):
        try:
            ba_values = pd.to_numeric(robust["BalancedAccuracyMean"], errors="coerce")
            if ba_values.notna().any():
                best_idx = ba_values.idxmax()
                row = robust.loc[best_idx]

                result["top_k"] = str(int(float(row["TopK"])))
                result["accuracy"] = fmt(row["AccuracyMean"])
                result["accuracy_std"] = fmt(row["AccuracyStd"])
                result["ba"] = fmt(row["BalancedAccuracyMean"])
                result["ba_std"] = fmt(row["BalancedAccuracyStd"])
                result["f1"] = fmt(row["F1MacroMean"])
                result["f1_std"] = fmt(row["F1MacroStd"])
                result["chance_ba"] = fmt(row["ChanceBalancedAccuracy"])

                result["status"] = (
                    "ABOVE CHANCE ONLY / NOT RELIABLE"
                    if normalize_bool(row["AboveChanceCandidate"])
                    else "NOT RELIABLE"
                )
        except Exception:
            pass

    return result


# ============================================================
# CONFUSION MATRIX
# ============================================================

def build_confusion_matrix(predictions: pd.DataFrame):
    required = {"ActualGroup", "PredictedGroup"}
    if predictions.empty or not required.issubset(predictions.columns):
        return pd.DataFrame()

    actual = predictions["ActualGroup"].astype(str)
    predicted = predictions["PredictedGroup"].astype(str)
    labels = sorted(set(actual) | set(predicted))

    matrix = pd.crosstab(
        actual,
        predicted,
        rownames=["Actual"],
        colnames=["Predicted"],
        dropna=False,
    )

    return matrix.reindex(
        index=labels,
        columns=labels,
        fill_value=0,
    )


# ============================================================
# ERROR ANALYSIS
# ============================================================

def build_error_analysis(predictions: pd.DataFrame):
    required = {"ActualGroup", "PredictedGroup"}
    if predictions.empty or not required.issubset(predictions.columns):
        return pd.DataFrame(), pd.DataFrame()

    df = predictions.copy()
    df["ActualGroup"] = df["ActualGroup"].astype(str)
    df["PredictedGroup"] = df["PredictedGroup"].astype(str)
    df["CorrectCalc"] = df["ActualGroup"] == df["PredictedGroup"]

    class_summary = (
        df.groupby("ActualGroup")
        .agg(
            Samples=("ActualGroup", "size"),
            Correct=("CorrectCalc", "sum"),
        )
        .reset_index()
    )
    class_summary["Errors"] = class_summary["Samples"] - class_summary["Correct"]
    class_summary["Accuracy"] = class_summary["Correct"] / class_summary["Samples"]
    class_summary["ErrorRate"] = 1 - class_summary["Accuracy"]

    class_summary = class_summary.sort_values(
        ["ErrorRate", "Samples"], ascending=[False, False]
    )

    errors = df[~df["CorrectCalc"]]
    if errors.empty:
        confusion_summary = pd.DataFrame()
    else:
        confusion_summary = (
            errors.groupby(["ActualGroup", "PredictedGroup"])
            .size()
            .reset_index(name="Errors")
            .sort_values("Errors", ascending=False)
        )

    return class_summary, confusion_summary


# ============================================================
# CONFIDENCE DISTRIBUTION
# ============================================================

def build_confidence_distribution(predictions: pd.DataFrame):
    if predictions.empty or "PredictionConfidence" not in predictions.columns:
        return pd.DataFrame()

    confidence = pd.to_numeric(
        predictions["PredictionConfidence"], errors="coerce"
    ).dropna()
    if confidence.empty:
        return pd.DataFrame()

    bins = [0.0, 0.1, 0.2, 0.3, 0.4, 0.5, 0.6, 0.7, 0.8, 0.9, 1.0]
    labels = [
        "0.0–0.1", "0.1–0.2", "0.2–0.3", "0.3–0.4", "0.4–0.5",
        "0.5–0.6", "0.6–0.7", "0.7–0.8", "0.8–0.9", "0.9–1.0",
    ]

    categories = pd.cut(
        confidence,
        bins=bins,
        labels=labels,
        include_lowest=True,
    )

    return categories.value_counts().sort_index().rename("Samples").to_frame()


# ============================================================
# HISTORY
# ============================================================

def save_current_run_history(robust: pd.DataFrame):
    if robust.empty:
        return
    if (
        "TopK" not in robust.columns
        or "BalancedAccuracyMean" not in robust.columns
        or "F1MacroMean" not in robust.columns
        or "AccuracyMean" not in robust.columns
    ):
        return
    try:
        HISTORY_DIR.mkdir(parents=True, exist_ok=True)
        timestamp = datetime.fromtimestamp(
            ROBUST_SUMMARY_PATH.stat().st_mtime
        ).strftime("%Y%m%d_%H%M%S")
        history_path = HISTORY_DIR / f"robust_{timestamp}.csv"
        if not history_path.exists():
            robust.to_csv(history_path, index=False, encoding="utf-8-sig")
    except Exception:
        pass


@st.cache_data(show_spinner=False)
def load_history():
    if not HISTORY_DIR.exists():
        return pd.DataFrame()

    files = sorted(HISTORY_DIR.glob("robust_*.csv"))
    rows = []

    for file in files:
        try:
            df = pd.read_csv(file)
            required = {"TopK", "AccuracyMean", "BalancedAccuracyMean", "F1MacroMean"}
            if not required.issubset(df.columns):
                continue
            ba = pd.to_numeric(df["BalancedAccuracyMean"], errors="coerce")
            if not ba.notna().any():
                continue
            row = df.loc[ba.idxmax()]
            rows.append(
                {
                    "Run": file.stem.replace("robust_", ""),
                    "TopK": int(float(row["TopK"])),
                    "Accuracy": float(row["AccuracyMean"]),
                    "BalancedAccuracy": float(row["BalancedAccuracyMean"]),
                    "MacroF1": float(row["F1MacroMean"]),
                }
            )
        except Exception:
            continue

    if not rows:
        return pd.DataFrame()
    return pd.DataFrame(rows).sort_values("Run")


# ============================================================
# LOAD CURRENT DATA
# ============================================================

robust_current = read_csv(str(ROBUST_SUMMARY_PATH))
save_current_run_history(robust_current)
history = load_history()
dataset = get_dataset_info()
metrics = get_metrics()
readme_text = read_text(str(README_PATH))
conclusion_text = read_text(str(CONCLUSION_PATH))
run_output = read_text(str(RUN_OUTPUT_PATH))
color_dashboard = read_json(str(COLOR_DASHBOARD_JSON))


# ============================================================
# HEADER
# ============================================================

st.title("🧪 داشبورد پژوهشی زعفران — NMR")
st.caption(
    "Intelligent System for Geographical Origin, Authenticity and Novelty Analysis Based on NMR Spectroscopic Data"
)
st.warning(
    f"⚠️ وضعیت عملیاتی مدل منشأ جغرافیایی: {metrics['status']}"
)

if COLOR_DASHBOARD_JSON.exists():
    color_samples_header = color_dashboard.get("samples", []) or []
    color_df_header = pd.DataFrame(color_samples_header)

    header_total = len(color_df_header)
    header_real = 0
    header_unknown = 0
    header_labeled_color = 0

    if not color_df_header.empty:
        if "knownArtificialColor" in color_df_header.columns:
            header_labeled_color = int(
                color_df_header["knownArtificialColor"]
                .apply(normalize_bool)
                .sum()
            )
        elif "class" in color_df_header.columns:
            class_series = color_df_header["class"].astype(str).str.lower()
            header_labeled_color = int(
                (~class_series.isin({"real_saffron", "unknown_adulterated_saffron"}))
                .sum()
            )

        if "class" in color_df_header.columns:
            class_series = color_df_header["class"].astype(str).str.lower()
            header_real = int((class_series == "real_saffron").sum())
            header_unknown = int((class_series == "unknown_adulterated_saffron").sum())

    st.caption(
        "🎨 تحلیل رنگ/اصالت Pilot فعال است: "
        f"{header_total} نمونه، "
        f"{header_labeled_color} نمونه دارای برچسب رنگ مصنوعی، "
        f"{header_real} نمونه مرجع زعفران واقعی و "
        f"{header_unknown} نمونه با تقلب ناشناخته. "
        "این نتایج هنوز اعتبارسنجی مستقل ندارند."
    )

# ============================================================
# TOP KPIs
# ============================================================

k1, k2, k3, k4, k5 = st.columns(5)
with k1:
    st.metric("نمونه‌های مستقل", metrics["samples"])
with k2:
    st.metric("گروه‌های جغرافیایی", metrics["groups"])
with k3:
    st.metric("نمونه‌های OOD", metrics["ood"])
with k4:
    st.metric("Robust Balanced Accuracy", metrics["ba"])
with k5:
    st.metric("Robust Macro F1", metrics["f1"])

st.divider()


# ============================================================
# SIDEBAR
# ============================================================

with st.sidebar:
    st.header("🧪 Saffron NMR")
    st.write("داشبورد جامع تحلیل، کنترل کیفیت و گزارش پژوهش NMR زعفران.")
    st.divider()
    st.subheader("Pipeline")

    pipeline = {
        "Dataset": DATASET_PATH,
        "PCA": PCA_SCORES_PLOT,
        "Harvest Year": YEAR_PCA_PLOT,
        "OOD": NOVELTY_PATH,
        "Robust Model": ROBUST_SUMMARY_PATH,
        "SHAP": SHAP_PATH,
        "Decision Engine": DECISION_PATH,
        "Final Report": FINAL_REPORT_PATH,
        "Region / Year Analysis": REGION_YEAR_SUMMARY_PATH,
        "Color / Authenticity": COLOR_DASHBOARD_JSON,
    }

    for name, path in pipeline.items():
        st.write(f"✅ {name}" if path.exists() else f"⚠️ {name}")

    st.divider()
    st.subheader("نمودارها")
    plots = [
        PCA_SCORES_PLOT,
        PCA_SCREE_PLOT,
        YEAR_PCA_PLOT,
        YEAR_REGION_PCA_PLOT,
        NOVELTY_PLOT,
        SHAP_PLOT,
    ]
    available_plots = sum(path.exists() for path in plots)
    st.write(f"{available_plots} از {len(plots)} آماده است.")

    st.divider()
    st.subheader("آخرین بروزرسانی")
    st.caption(f"Dataset: {file_time(DATASET_PATH)}")
    st.caption(f"Robust: {file_time(ROBUST_SUMMARY_PATH)}")
    st.caption(f"Run: {file_time(RUN_OUTPUT_PATH)}")
    if REGION_YEAR_SUMMARY_PATH.exists():
        st.caption(f"Region/Year: {file_time(REGION_YEAR_SUMMARY_PATH)}")
    if COLOR_DASHBOARD_JSON.exists():
        st.caption(f"Color: {file_time(COLOR_DASHBOARD_JSON)}")

    st.divider()
    if st.button("🔄 بروزرسانی داشبورد", width="stretch"):
        st.cache_data.clear()
        st.rerun()


# ============================================================
# TABS
# ============================================================

(
    tab_summary,
    tab_qc,
    tab_pca,
    tab_origin,
    tab_cross_year,
    tab_shap,
    tab_sample,
    tab_color,
    tab_scientific,
    tab_run,
    tab_report,
    tab_readme,
    tab_conclusion,
    tab_region_year,
    tab_reference_doc,
) = st.tabs(
    [
        "🏠 خلاصه پژوهش",
        "🧪 کیفیت داده",
        "📈 PCA و سال برداشت",
        "🌍 مدل منشأ",
        "🔍 Cross-Year و OOD",
        "🧠 تفسیر مدل",
        "🔬 Sample Explorer",
        "🎨 تشخیص رنگ و اصالت",
        "📑 تفسیر علمی",
        "📋 Run Log",
        "📄 Final Report",
        "📘 README",
        "📝 نتیجه‌گیری",
        "🌍 تحلیل تخصصی منطقه و سال",
        "📘 سند مرجع پروژه",
    ]
)


# ============================================================
# TAB 1 — SUMMARY
# ============================================================

with tab_summary:
    st.subheader("نمای کلی پژوهش")
    left, right = st.columns(2)

    with left:
        st.markdown("### وضعیت داده")
        st.write(f"رکوردهای اولیه: **{dataset['initial_records']}**")
        st.write(f"نمونه‌های مستقل: **{dataset['independent_samples']}**")
        st.write(f"نقاط طیفی: **{dataset['spectral_points']}**")
        st.write(f"گروه‌های جغرافیایی: **{dataset['groups']}**")
        st.write(
            f"سال‌های موجود: **{', '.join(map(str, dataset['years'])) or '—'}**"
        )
        st.write(f"Missing Values: **{dataset['missing_values']}**")

    with right:
        st.markdown("### وضعیت مدل")
        st.write(f"مدل منتخب: **XGBoost / TopK {metrics['top_k']}**")
        st.write(f"Accuracy: **{metrics['accuracy']} ± {metrics['accuracy_std']}**")
        st.write(f"Balanced Accuracy: **{metrics['ba']} ± {metrics['ba_std']}**")
        st.write(f"Macro F1: **{metrics['f1']} ± {metrics['f1_std']}**")
        st.write(f"Chance BA: **{metrics['chance_ba']}**")
        st.write(f"OOD Candidates: **{metrics['ood']}**")

    st.divider()
    st.markdown("### وضعیت مراحل پژوهش")
    stage_df = pd.DataFrame(
        [
            ["کنترل کیفیت داده", "✅"],
            ["آماده‌سازی داده", "✅"],
            ["تحلیل سال برداشت", "✅"],
            ["PCA", "✅"],
            ["Novelty / OOD", "✅"],
            ["Feature Engineering", "✅"],
            ["Origin Modeling", "⚠️"],
            ["Robust Evaluation", "✅"],
            ["SHAP", "✅"],
            ["Decision Engine", "✅"],
            ["Final Report", "✅"],
            ["تشخیص رنگ مصنوعی / اصالت", "✅" if COLOR_DASHBOARD_JSON.exists() else "⏳"],
            ["تحلیل تخصصی منطقه و سال", "✅" if REGION_YEAR_SUMMARY_PATH.exists() else "⏳"],
        ],
        columns=["Component", "Status"],
    )
    st.dataframe(stage_df, hide_index=True, width="stretch")
    st.divider()
    st.info(
        "چارچوب محاسباتی پروژه اجرا شده است، اما مدل منشأ جغرافیایی در شرایط فعلی برای استفاده عملیاتی قابل اتکا نیست."
    )


# ============================================================
# TAB 2 — DATA QUALITY
# ============================================================

with tab_qc:
    st.subheader("کیفیت داده و کنترل کیفیت")
    q1, q2, q3, q4, q5 = st.columns(5)
    with q1:
        st.metric("رکوردهای اولیه", dataset["initial_records"])
    with q2:
        st.metric("نمونه‌های مستقل", dataset["independent_samples"])
    with q3:
        st.metric("Spectral Points", dataset["spectral_points"])
    with q4:
        st.metric("Missing Values", dataset["missing_values"])
    with q5:
        st.metric("Invalid Values", dataset["invalid_values"])

    st.divider()
    c1, c2 = st.columns(2)
    with c1:
        st.markdown("### توزیع گروه‌ها")
        groups = dataset["group_counts"]
        if not groups.empty:
            st.bar_chart(groups.set_index("Group")["Samples"])
            st.dataframe(groups, hide_index=True, width="stretch")
            st.caption(
                f"حداقل: {groups['Samples'].min()} | حداکثر: {groups['Samples'].max()}"
            )
    with c2:
        st.markdown("### توزیع سال برداشت")
        years = dataset["year_counts"]
        if not years.empty:
            st.bar_chart(years.set_index("HarvestYear")["Samples"])
            st.dataframe(years, hide_index=True, width="stretch")

    st.divider()
    st.markdown("### Region × Harvest Year")
    region_year = dataset["region_year"]
    if not region_year.empty:
        styled = region_year.style.background_gradient(cmap="Greens").format("{:.0f}")
        st.dataframe(styled, width="stretch")
        common_regions = int(sum(((region_year > 0).sum(axis=1) >= 2)))
        st.info(
            f"مناطق دارای نمونه در هر دو سال: **{common_regions} / {len(region_year)}**"
        )

    st.divider()
    st.markdown("### کنترل‌های ساختاری")
    qc_df = pd.DataFrame(
        [
            [
                "Missing Values",
                dataset["missing_values"],
                "PASS" if dataset["missing_values"] == 0 else "CHECK",
            ],
            [
                "Invalid Spectral Values",
                dataset["invalid_values"],
                "PASS" if dataset["invalid_values"] == 0 else "CHECK",
            ],
            [
                "Duplicate Extra Rows",
                dataset["duplicate_extra_rows"],
                "REMOVED" if dataset["duplicate_extra_rows"] > 0 else "PASS",
            ],
            [
                "Spectral Points",
                dataset["spectral_points"],
                "AVAILABLE" if dataset["spectral_points"] > 0 else "CHECK",
            ],
        ],
        columns=["Check", "Value", "Status"],
    )
    st.dataframe(qc_df, hide_index=True, width="stretch")
    if dataset["min_ppm"] is not None and dataset["max_ppm"] is not None:
        st.write(
            f"Spectral axis: **{dataset['min_ppm']:.6f} → {dataset['max_ppm']:.6f} ppm**"
        )


# ============================================================
# TAB 3 — PCA / YEAR
# ============================================================

with tab_pca:
    st.subheader("PCA و تحلیل سال برداشت")
    pca1, pca2 = st.columns(2)
    with pca1:
        show_image(PCA_SCORES_PLOT, "PCA Scores")
    with pca2:
        show_image(PCA_SCREE_PLOT, "PCA Scree Plot")
    st.divider()
    st.markdown("### تحلیل سال برداشت")
    y1, y2 = st.columns(2)
    with y1:
        show_image(YEAR_PCA_PLOT, "Harvest Year PCA")
    with y2:
        show_image(YEAR_REGION_PCA_PLOT, "Region × Harvest Year PCA")
    st.divider()

    year_prediction = read_csv(str(YEAR_PREDICTION_PATH))
    if not year_prediction.empty:
        st.markdown("### Cross-Validated Year Prediction")
        st.dataframe(year_prediction, hide_index=True, width="stretch")

    year_effect = read_csv(str(YEAR_EFFECT_PATH))
    if not year_effect.empty:
        st.markdown("### Within-Region Year Effect")
        st.dataframe(year_effect, hide_index=True, width="stretch")

    year_summary = read_csv(str(YEAR_EFFECT_SUMMARY_PATH))
    if not year_summary.empty:
        st.markdown("### Year Effect Summary")
        st.dataframe(year_summary, hide_index=True, width="stretch")


# ============================================================
# TAB 4 — ORIGIN MODEL
# ============================================================

with tab_origin:
    st.subheader("مدل‌سازی منشأ جغرافیایی")
    robust = read_csv(str(ROBUST_SUMMARY_PATH))

    if robust.empty:
        st.warning("robust_evaluation_summary.csv موجود نیست.")
    else:
        required = {
            "TopK", "AccuracyMean", "AccuracyStd", "BalancedAccuracyMean",
            "BalancedAccuracyStd", "F1MacroMean", "F1MacroStd",
        }
        if required.issubset(robust.columns):
            table = robust.copy()
            table["Accuracy"] = table.apply(
                lambda row: fmt_pm(row["AccuracyMean"], row["AccuracyStd"]), axis=1
            )
            table["Balanced Accuracy"] = table.apply(
                lambda row: fmt_pm(
                    row["BalancedAccuracyMean"], row["BalancedAccuracyStd"]
                ), axis=1
            )
            table["Macro F1"] = table.apply(
                lambda row: fmt_pm(row["F1MacroMean"], row["F1MacroStd"]), axis=1
            )
            table = table[["TopK", "Accuracy", "Balanced Accuracy", "Macro F1"]]
            st.dataframe(table, hide_index=True, width="stretch")

            ba_chart = robust[["TopK", "BalancedAccuracyMean"]].copy()
            ba_chart["TopK"] = ba_chart["TopK"].astype(int).astype(str)
            st.markdown("### Balanced Accuracy")
            st.bar_chart(ba_chart.set_index("TopK")["BalancedAccuracyMean"])

            f1_chart = robust[["TopK", "F1MacroMean"]].copy()
            f1_chart["TopK"] = f1_chart["TopK"].astype(int).astype(str)
            st.markdown("### Macro F1")
            st.bar_chart(f1_chart.set_index("TopK")["F1MacroMean"])

    st.divider()
    st.markdown("### Confusion Matrix")
    predictions = read_csv(str(PREDICTIONS_PATH))
    confusion = build_confusion_matrix(predictions)
    if confusion.empty:
        st.info("Confusion Matrix قابل محاسبه نیست.")
    else:
        st.dataframe(
            confusion.style.background_gradient(cmap="Greens"),
            width="stretch",
        )

    st.divider()
    st.markdown("### Error Analysis")
    class_errors, confusion_errors = build_error_analysis(predictions)
    if not class_errors.empty:
        display_errors = class_errors.copy()
        display_errors["Accuracy"] = display_errors["Accuracy"].map(fmt)
        display_errors["ErrorRate"] = display_errors["ErrorRate"].map(fmt)
        st.dataframe(display_errors, hide_index=True, width="stretch")

    if not confusion_errors.empty:
        st.markdown("#### پرتکرارترین Confusionها")
        st.dataframe(confusion_errors.head(15), hide_index=True, width="stretch")

    st.divider()
    st.markdown("### Prediction Confidence")
    confidence_dist = build_confidence_distribution(predictions)
    if confidence_dist.empty:
        st.info("اطلاعات Confidence موجود نیست.")
    else:
        st.bar_chart(confidence_dist["Samples"])
        confidence_values = pd.to_numeric(
            predictions["PredictionConfidence"], errors="coerce"
        ).dropna()
        if not confidence_values.empty:
            c1, c2, c3 = st.columns(3)
            with c1:
                st.metric("Mean Confidence", fmt(confidence_values.mean()))
            with c2:
                st.metric("Maximum Confidence", fmt(confidence_values.max()))
            with c3:
                st.metric("Minimum Confidence", fmt(confidence_values.min()))

    st.divider()
    st.warning(
        f"مدل منتخب فعلی **XGBoost / TopK {metrics['top_k']}** است، اما وضعیت عملیاتی آن **{metrics['status']}** است."
    )


# ============================================================
# TAB 5 — CROSS YEAR / OOD
# ============================================================

with tab_cross_year:
    st.subheader("Cross-Year Generalization و Novelty / OOD")
    st.markdown("### Cross-Year")
    cross_year = read_csv(str(CROSS_YEAR_PATH))
    if not cross_year.empty:
        st.dataframe(cross_year, hide_index=True, width="stretch")
        if "BalancedAccuracy" in cross_year.columns:
            best = cross_year.sort_values("BalancedAccuracy", ascending=False).head(1)
            if not best.empty:
                row = best.iloc[0]
                st.info(
                    f"""
                    بهترین Cross-Year Candidate:

                    **{int(row['SourceYear'])} → {int(row['TargetYear'])}**

                    Model: **{row['Model']}**

                    TopK: **{int(row['TopK'])}**

                    Balanced Accuracy: **{fmt(row['BalancedAccuracy'])}**

                    این نتیجه یک تست تعمیم بین‌سال است و عملکرد عملیاتی نهایی محسوب نمی‌شود.
                    """
                )
    else:
        st.info("نتایج Cross-Year موجود نیست.")

    st.divider()
    st.markdown("### Novelty / OOD")
    novelty = read_csv(str(NOVELTY_PATH))
    ood1, ood2 = st.columns(2)
    with ood1:
        st.metric("OOD Candidates", metrics["ood"])
        if not novelty.empty:
            status_column = find_column(novelty, ["NoveltyStatus", "Status"])
            if status_column:
                novel_rows = novelty[
                    novelty[status_column]
                    .astype(str)
                    .str.contains(r"NOVEL|OOD", case=False, regex=True)
                ]
            else:
                novel_rows = novelty
            if not novel_rows.empty:
                st.dataframe(novel_rows, hide_index=True, width="stretch")
    with ood2:
        show_image(NOVELTY_PLOT, "Mahalanobis Novelty Detection")
    st.warning(
        "OOD ≠ Adulteration\n\nخارج بودن یک نمونه از دامنه مرجع به‌تنهایی تقلب یا ناخالصی را اثبات نمی‌کند."
    )


# ============================================================
# TAB 6 — SHAP
# ============================================================

with tab_shap:
    st.subheader("تفسیر مدل و SHAP")
    s1, s2 = st.columns(2)
    with s1:
        show_image(SHAP_PLOT, "SHAP Summary Plot")
    with s2:
        shap_df = read_csv(str(SHAP_PATH))
        if not shap_df.empty:
            columns = [
                col for col in ["Rank", "Feature", "MeanAbsSHAP"]
                if col in shap_df.columns
            ]
            st.markdown("### Top Spectral Features")
            st.dataframe(
                shap_df[columns].head(20),
                hide_index=True,
                width="stretch",
            )
    st.divider()
    st.info(
        "ویژگی‌های SHAP نواحی طیفی مرتبط با تصمیم مدل را نشان می‌دهند. این نواحی biomarkerهای تأییدشده منشأ جغرافیایی محسوب نمی‌شوند."
    )


# ============================================================
# TAB 7 — SAMPLE EXPLORER
# ============================================================

with tab_sample:
    st.subheader("Sample Explorer")
    predictions = read_csv(str(PREDICTIONS_PATH))
    novelty = read_csv(str(NOVELTY_PATH))
    decision = read_csv(str(DECISION_PATH))

    if predictions.empty:
        st.warning("sample_predictions.csv موجود نیست.")
    elif "SampleId" not in predictions.columns:
        st.warning("ستون SampleId موجود نیست.")
    else:
        sample_ids = predictions["SampleId"].dropna().astype(int).tolist()
        selected_id = st.selectbox("انتخاب نمونه", sample_ids)
        selected_df = predictions[predictions["SampleId"] == selected_id]

        if not selected_df.empty:
            row = selected_df.iloc[0]
            c1, c2, c3 = st.columns(3)
            with c1:
                st.metric("Sample ID", selected_id)
                if "ActualGroup" in row:
                    st.write(f"گروه واقعی: **{row['ActualGroup']}**")
            with c2:
                if "PredictedGroup" in row:
                    st.write(f"پیش‌بینی مدل: **{row['PredictedGroup']}**")
                if "PredictionConfidence" in row:
                    confidence = float(row["PredictionConfidence"])
                    st.write(f"Confidence: **{confidence:.4f}**")
            with c3:
                if "Correct" in row:
                    correct = normalize_bool(row["Correct"])
                    st.write(f"پیش‌بینی صحیح: **{'بله' if correct else 'خیر'}**")

            st.divider()
            sample_ood = pd.DataFrame()
            if not novelty.empty and "SampleId" in novelty.columns:
                sample_ood = novelty[novelty["SampleId"] == selected_id]
            if not sample_ood.empty:
                st.error("🔍 این نمونه OOD Candidate است.")
                st.markdown("### جزئیات OOD")
                st.dataframe(sample_ood, hide_index=True, width="stretch")
            else:
                st.success("این نمونه در دامنه مرجع قرار دارد.")

            if "PredictionConfidence" in row:
                confidence = float(row["PredictionConfidence"])
                st.markdown("### اطمینان پیش‌بینی")
                st.progress(min(max(confidence, 0.0), 1.0))

            if not decision.empty and "SampleId" in decision.columns:
                sample_decision = decision[decision["SampleId"] == selected_id]
                if not sample_decision.empty:
                    st.divider()
                    st.markdown("### Decision Engine")
                    st.dataframe(sample_decision, hide_index=True, width="stretch")


# ============================================================
# TAB 8 — ARTIFICIAL COLOR / AUTHENTICITY
# ============================================================

with tab_color:
    st.subheader("🎨 تشخیص رنگ مصنوعی و بررسی اصالت")

    if not color_dashboard:
        st.warning("نتایج تحلیل color.csv هنوز تولید نشده‌اند.")
        st.code("python -m src.color_dashboard_result", language="powershell")
    else:
        analysis = color_dashboard.get("analysis", {}) or {}
        threshold = color_dashboard.get("threshold", {}) or {}
        summary = color_dashboard.get("summary", {}) or {}
        samples = color_dashboard.get("samples", []) or []
        region = analysis.get("region", {}) or {}

        sample_df = pd.DataFrame(samples)

        # ----------------------------------------------------
        # Reference / label counts
        # ----------------------------------------------------
        total_samples = (
            len(sample_df)
            if not sample_df.empty
            else summary.get("totalSamples", 0)
        )

        real_reference_count = 0
        unknown_adulteration_count = 0
        labeled_artificial_color_count = 0

        if not sample_df.empty:
            if "class" in sample_df.columns:
                class_series = sample_df["class"].astype(str).str.lower()
                real_reference_count = int(
                    (class_series == "real_saffron").sum()
                )
                unknown_adulteration_count = int(
                    (class_series == "unknown_adulterated_saffron").sum()
                )

            if "knownArtificialColor" in sample_df.columns:
                labeled_artificial_color_count = int(
                    sample_df["knownArtificialColor"]
                    .apply(normalize_bool)
                    .sum()
                )
            elif "class" in sample_df.columns:
                labeled_artificial_color_count = int(
                    (~class_series.isin(
                        {
                            "real_saffron",
                            "unknown_adulterated_saffron",
                        }
                    )).sum()
                )

        threshold_value = threshold.get("value", None)
        balanced_accuracy = threshold.get("balancedAccuracy", None)
        sensitivity = threshold.get("sensitivity", None)
        specificity = threshold.get("specificity", None)

        if balanced_accuracy is None:
            balanced_accuracy = analysis.get("balancedAccuracy", None)
        if sensitivity is None:
            sensitivity = analysis.get("sensitivity", None)
        if specificity is None:
            specificity = analysis.get("specificity", None)

        st.caption(
            f"ناحیه مورد استفاده: {region.get('lowerPpm', 5.0)} تا "
            f"{region.get('upperPpm', 9.0)} ppm"
        )

        st.warning(
            "این بخش بر اساس داده Pilot فعلی و یک روش اکتشافی ساخته شده است و هنوز برای استفاده تولیدی یا تشخیص مستقل اعتبارسنجی نشده است."
        )

        st.info(
            "نکته مهم: تعداد «نمونه‌های دارای برچسب رنگ مصنوعی» از برچسب مرجع Dataset گزارش می‌شود و به معنی کشف مستقل این تعداد نمونه توسط مدل نیست."
        )

        # ----------------------------------------------------
        # KPI
        # ----------------------------------------------------

        c1, c2, c3, c4, c5 = st.columns(5)

        with c1:
            st.metric("کل نمونه‌ها", total_samples)

        with c2:
            st.metric(
                "نمونه‌های دارای برچسب رنگ مصنوعی",
                labeled_artificial_color_count,
            )

        with c3:
            st.metric(
                "نمونه‌های مرجع زعفران واقعی",
                real_reference_count,
            )

        with c4:
            st.metric(
                "تقلب ناشناخته",
                unknown_adulteration_count,
            )

        with c5:
            st.metric(
                "آستانه اکتشافی Pilot",
                fmt(threshold_value, 4)
                if threshold_value is not None
                else "—",
            )

        st.divider()

        # ----------------------------------------------------
        # Spectral focus
        # ----------------------------------------------------

        st.markdown("### ناحیه طیفی تحلیل رنگ")

        r1, r2, r3 = st.columns(3)

        with r1:
            st.metric("شروع ناحیه", f"{region.get('lowerPpm', 5.0):.3f} ppm")

        with r2:
            st.metric("پایان ناحیه", f"{region.get('upperPpm', 9.0):.3f} ppm")

        with r3:
            st.metric("ناحیه حلال حذف‌شده", "4.66797–5.4996 ppm")

        st.info(
            "بر اساس طراحی آزمایش، تفسیر شواهد رنگ مصنوعی فقط در محدوده 5 تا 9 ppm انجام می‌شود. ناحیه حلال 4.66797 تا 5.4996 ppm از ماتریس تحلیل حذف شده است."
        )

        # ----------------------------------------------------
        # Exploratory validation
        # ----------------------------------------------------

        st.markdown("### ارزیابی داخلی اکتشافی")

        v1, v2, v3 = st.columns(3)

        with v1:
            st.metric(
                "Balanced Accuracy",
                fmt(balanced_accuracy, 4)
                if balanced_accuracy is not None
                else "—",
            )

        with v2:
            st.metric(
                "Sensitivity",
                fmt(sensitivity, 4)
                if sensitivity is not None
                else "—",
            )

        with v3:
            st.metric(
                "Specificity",
                fmt(specificity, 4)
                if specificity is not None
                else "—",
            )

        st.caption(
            "این شاخص‌ها از Leave-One-Out روی همین Pilot Dataset به‌دست آمده‌اند. Threshold نیز با همین Pilot انتخاب شده و بنابراین این اعداد Validation مستقل محسوب نمی‌شوند."
        )

        # ----------------------------------------------------
        # Exploratory plots
        # ----------------------------------------------------

        if COLOR_EDA_REGION_PLOT.exists() or COLOR_PCA_REGION_PLOT.exists():
            st.divider()
            st.markdown("### نمودارهای اکتشافی")

            if COLOR_EDA_REGION_PLOT.exists() and COLOR_PCA_REGION_PLOT.exists():
                p1, p2 = st.columns(2)
                with p1:
                    show_image(
                        COLOR_EDA_REGION_PLOT,
                        "Color Region 5–9 ppm",
                    )
                with p2:
                    show_image(
                        COLOR_PCA_REGION_PLOT,
                        "PCA — Color Region",
                    )
            elif COLOR_EDA_REGION_PLOT.exists():
                show_image(
                    COLOR_EDA_REGION_PLOT,
                    "Color Region 5–9 ppm",
                )
            elif COLOR_PCA_REGION_PLOT.exists():
                show_image(
                    COLOR_PCA_REGION_PLOT,
                    "PCA — Color Region",
                )

        # ----------------------------------------------------
        # Decision distribution
        # ----------------------------------------------------

        st.divider()
        st.markdown("### توزیع تصمیم‌های اکتشافی")

        decision_counts = summary.get("decisionCounts", {}) or {}

        if decision_counts:
            decision_df = pd.DataFrame(
                [
                    {"Decision": key, "Samples": value}
                    for key, value in decision_counts.items()
                ]
            )

            st.dataframe(
                decision_df,
                hide_index=True,
                width="stretch",
            )

            st.bar_chart(
                decision_df.set_index("Decision")["Samples"]
            )

        st.info(
            "این تصمیم‌ها باید در کنار برچسب مرجع و محدودیت‌های Pilot تفسیر شوند و به‌تنهایی به معنی اثبات تقلب نیستند."
        )

        # ----------------------------------------------------
        # Sample table
        # ----------------------------------------------------

        st.divider()
        st.markdown("### نتایج نمونه‌ها")

        if sample_df.empty:
            st.info("نتیجه‌ای برای نمونه‌ها موجود نیست.")
        else:
            display_columns = [
                "number",
                "name",
                "class",
                "decision",
                "status",
                "score",
                "confidence",
                "artificialColorPercent",
            ]

            available_columns = [
                c for c in display_columns
                if c in sample_df.columns
            ]

            display_df = sample_df[available_columns].copy()

            display_df = display_df.rename(
                columns={
                    "number": "شماره",
                    "name": "نمونه",
                    "class": "کلاس مرجع",
                    "decision": "تصمیم",
                    "status": "وضعیت",
                    "score": "امتیاز اکتشافی",
                    "confidence": "اعتماد اکتشافی",
                    "artificialColorPercent": "درصد رنگ مصنوعی ثبت‌شده",
                }
            )

            st.dataframe(
                display_df,
                hide_index=True,
                width="stretch",
            )

        # ----------------------------------------------------
        # Sample detail
        # ----------------------------------------------------

        st.divider()
        st.markdown("### بررسی یک نمونه")

        if not sample_df.empty and "number" in sample_df.columns:
            selected_number = st.selectbox(
                "نمونه را انتخاب کنید",
                sample_df["number"].tolist(),
                format_func=lambda value: (
                    f"#{int(value)} — "
                    f"{sample_df.loc[sample_df['number'] == value, 'name'].iloc[0]}"
                    if "name" in sample_df.columns
                    and not sample_df.loc[
                        sample_df["number"] == value
                    ].empty
                    else f"#{int(value)}"
                ),
                key="color_sample_selector",
            )

            selected_rows = sample_df[
                sample_df["number"] == selected_number
            ]

            if not selected_rows.empty:
                selected = selected_rows.iloc[0]

                a, b, c, d = st.columns(4)

                with a:
                    st.metric("Sample", int(selected_number))

                with b:
                    st.metric(
                        "امتیاز اکتشافی",
                        fmt(selected.get("score", None), 4),
                    )

                with c:
                    st.metric(
                        "وضعیت",
                        str(selected.get("status", "—")),
                    )

                with d:
                    st.metric(
                        "تصمیم",
                        str(selected.get("decision", "—")),
                    )

                st.markdown(
                    f"""
                    **نام نمونه:** {selected.get('name', '—')}

                    **کلاس مرجع:** {selected.get('class', '—')}

                    **دلیل:** {selected.get('reason', '—')}
                    """
                )

                artificial_percent = selected.get(
                    "artificialColorPercent",
                    None,
                )

                if pd.notna(artificial_percent):
                    st.write(
                        "درصد رنگ مصنوعی ثبت‌شده در Metadata: "
                        f"**{float(artificial_percent):.1f}%**"
                    )

        # ----------------------------------------------------
        # Scientific limitations
        # ----------------------------------------------------

        st.divider()
        st.markdown("### محدودیت‌های علمی")

        limitations = color_dashboard.get("limitations", []) or []

        if limitations:
            for limitation in limitations:
                st.warning(limitation)
        else:
            st.warning(
                "مرجع زعفران واقعی Pilot محدود است؛ Threshold بر اساس همین Pilot تعیین شده و Validation مستقل وجود ندارد."
            )

        st.error(
            "نمونه‌های دارای تقلب ناشناخته مانند 650 و 651 نباید صرفاً بر اساس این نمره رنگی قضاوت شوند و باید در چارچوب Whole-Spectrum/OOD بررسی شوند."
        )

        st.divider()

        if COLOR_DASHBOARD_CSV.exists():
            st.download_button(
                label="⬇️ دانلود نتایج تشخیص رنگ",
                data=COLOR_DASHBOARD_CSV.read_bytes(),
                file_name="color_dashboard_samples.csv",
                mime="text/csv",
                width="content",
            )


# TAB 9 — SCIENTIFIC INTERPRETATION
# ============================================================

with tab_scientific:
    st.subheader("Scientific Interpretation")

    st.markdown("### 🟢 یافته‌های پشتیبانی‌شده")

    st.success(
        "کنترل کیفیت، آماده‌سازی داده، PCA، Novelty/OOD، ارزیابی مدل و تفسیر SHAP در Pipeline اجرا شده‌اند."
    )

    st.markdown("### 🟡 یافته‌های قابل بررسی")

    st.warning(
        "ساختار مرتبط با سال برداشت در داده مشاهده شده است، اما آزمون کنترل‌شده درون‌منطقه‌ای به معنی‌داری آماری نرسیده است."
    )

    st.warning(
        "یک یا چند نمونه ممکن است نسبت به دامنه مرجع فعلی Novel/OOD باشند، اما این امر به‌تنهایی تقلب را ثابت نمی‌کند."
    )

    st.markdown("### 🔴 یافته‌های تأییدنشده")

    st.error(
        "مدل منشأ جغرافیایی فعلی برای استفاده عملیاتی قابل اعتماد نیست."
    )

    st.markdown("### 🎨 Authenticity / Artificial Color")

    st.warning(
        """
        تحلیل Pilot داده‌های مربوط به رنگ‌های مصنوعی نشان می‌دهد که
        بررسی محدوده طیفی 5 تا 9 ppm می‌تواند یک مسیر بالقوه برای
        شناسایی شواهد مرتبط با افزودن رنگ مصنوعی به زعفران باشد.

        با این حال، تعداد نمونه‌های زعفران واقعی مرجع در این Pilot
        بسیار محدود است، Threshold بر اساس همین مجموعه اکتشافی تعیین
        شده و Validation مستقل انجام نشده است. بنابراین شاخص‌های
        فعلی را نمی‌توان به‌عنوان عملکرد قطعی یا عملیاتی یک مدل عمومی
        تشخیص رنگ مصنوعی تعمیم داد.

        همچنین ممکن است مقادیر بسیار کم رنگ مصنوعی به دلیل محدودیت
        حساسیت NMR قابل تشخیص نباشند. در نتیجه، این بخش در وضعیت
        فعلی یک Proof-of-Concept پژوهشی محسوب می‌شود.
        """
    )

    st.divider()

    st.markdown("### محدودیت‌های اصلی")

    st.write(
        f"""
        نمونه‌های مستقل فعلی: **{metrics['samples']}**

        گروه‌های جغرافیایی: **{metrics['groups']}**

        عدم توازن کلاس‌ها، حجم کم برخی مناطق، پوشش ناقص منطقه × سال
        و نبود مجموعه مستقل و معتبر برای اصالت و تقلب، از محدودیت‌های
        اصلی مطالعه فعلی هستند.
        """
    )

    st.divider()

    st.markdown("### وضعیت کلی پژوهش")

    st.info(
        """
        **Research Proof-of-Concept**

        چارچوب محاسباتی پروژه از نظر پردازش داده، تحلیل اثر سال،
        مدل‌سازی منشأ، Novelty/OOD، تفسیر مدل و زیرساخت اولیه
        ارزیابی اصالت قابل اجرا و بازتولید است؛ اما حجم و طراحی
        فعلی Dataset برای ارائه یک مدل قطعی و عملیاتی کافی نیست.
        """
    )


# TAB 10 — RUN LOG
# ============================================================

with tab_run:
    st.subheader("Pipeline Run Log")
    st.caption(f"آخرین بروزرسانی: {file_time(RUN_OUTPUT_PATH)}")
    if not run_output:
        st.warning("run_output.txt موجود نیست.")
    else:
        stages = [
            ("Validation", "VALIDATION FINISHED"),
            ("Preparation", "PREPARATION COMPLETED"),
            ("Harvest Year", "HARVEST YEAR ANALYSIS FINISHED"),
            ("PCA", "PCA ANALYSIS FINISHED"),
            ("Novelty / OOD", "MAHALANOBIS NOVELTY / OOD ANALYSIS FINISHED"),
            ("Feature Engineering", "FEATURE ENGINEERING FINISHED"),
            ("Origin Prediction", "ORIGIN PREDICTION FINISHED"),
            ("Robust Evaluation", "ROBUST MODEL EVALUATION FINISHED"),
            ("SHAP", "SHAP ANALYSIS FINISHED"),
            ("Decision Engine", "DECISION ENGINE FINISHED"),
            ("Final Report", "FINAL RESEARCH REPORT FINISHED"),
            ("Main Pipeline", "MAIN PIPELINE FINISHED"),
        ]
        run_df = pd.DataFrame(
            [[name, "✅" if marker in run_output else "—"] for name, marker in stages],
            columns=["Stage", "Status"],
        )
        st.dataframe(run_df, hide_index=True, width="stretch")
        completed = int((run_df["Status"] == "✅").sum())
        total = len(run_df)
        st.metric("Pipeline Completion", f"{completed} / {total}")
        st.progress(completed / total if total else 0)
        st.divider()
        with st.expander("▶ نمایش کامل run_output.txt"):
            st.code(run_output, language="text")


# ============================================================
# TAB 11 — FINAL REPORT
# ============================================================

with tab_report:
    st.subheader("Final Research Report")
    final_report = read_csv(str(FINAL_REPORT_PATH))
    if final_report.empty:
        st.warning("final_report.csv موجود نیست یا خالی است.")
    else:
        st.dataframe(final_report, hide_index=True, width="stretch")
        st.download_button(
            label="⬇️ دانلود Final Report",
            data=final_report.to_csv(index=False).encode("utf-8-sig"),
            file_name="final_report.csv",
            mime="text/csv",
        )


# ============================================================
# TAB 12 — README
# ============================================================

with tab_readme:
    st.subheader("Project Documentation")
    if not readme_text.strip():
        st.warning("README.md موجود نیست یا خالی است.")
    else:
        st.markdown(readme_text)


# ============================================================
# TAB 13 — CONCLUSION
# ============================================================

with tab_conclusion:
    st.subheader("نتیجه‌گیری علمی پژوهش")
    if not conclusion_text.strip():
        st.warning("docs/conclusion_fa.md موجود نیست یا خالی است.")
    else:
        st.markdown(conclusion_text)

        st.divider()

        st.markdown("### 🎨 نتیجه‌گیری درباره اصالت و رنگ مصنوعی")

        st.warning(
            """
            تحلیل Pilot مربوط به رنگ‌های مصنوعی نشان داد که بررسی
            محدوده طیفی 5 تا 9 ppm می‌تواند یک مسیر بالقوه برای
            شناسایی شواهد مرتبط با افزودن رنگ مصنوعی به زعفران باشد.

            با این حال، تعداد نمونه‌های زعفران واقعی مرجع در Pilot
            فعلی محدود است و Threshold نیز بر اساس همین مجموعه
            اکتشافی تعیین شده است. در نتیجه، شاخص‌های به‌دست‌آمده
            عملکرد مستقل یک مدل در برابر نمونه‌های جدید را نشان
            نمی‌دهند و برای تأیید قطعی اصالت یا اثبات قطعی تقلب کافی نیستند.

            علاوه بر این، مقادیر بسیار کم رنگ مصنوعی ممکن است به دلیل
            محدودیت حساسیت NMR قابل شناسایی نباشند. بنابراین، این
            بخش در وضعیت فعلی باید به‌عنوان یک Proof-of-Concept در نظر
            گرفته شود و توسعه مدل عملیاتی آن نیازمند نمونه‌های مرجع
            بیشتر و Validation مستقل است.
            """
        )

        if COLOR_DASHBOARD_JSON.exists():
            color_summary = color_dashboard.get("summary", {}) or {}
            color_threshold = color_dashboard.get("threshold", {}) or {}
            color_counts = color_summary.get("statusCounts", {}) or {}

            st.caption(
                "خلاصه Pilot: "
                f"{color_summary.get('totalSamples', 0)} نمونه، "
                f"{color_counts.get('positive', 0)} نمونه دارای نتیجه مثبت، "
                f"{color_counts.get('suspicious', 0)} نمونه مشکوک و "
                f"{color_counts.get('reference', 0)} نمونه مرجع."
            )

            st.caption(
                "اعداد بالا توصیف مجموعه Pilot هستند و به معنی عملکرد اعتبارسنجی‌شده در نمونه‌های مستقل جدید نیستند."
            )

    st.divider()
    st.markdown("### وضعیت کلی پژوهش")
    st.info(
        """
        **Research Proof-of-Concept**

        چارچوب محاسباتی پروژه از نظر پردازش داده، تحلیل اثر سال،
        مدل‌سازی منشأ، Novelty/OOD، تفسیر مدل و زیرساخت اولیه
        ارزیابی اصالت قابل اجرا و بازتولید است؛ اما حجم و طراحی
        فعلی Dataset برای ارائه یک مدل قطعی و عملیاتی کافی نیست.

        بنابراین مسیر اصلی فاز بعدی، افزایش نمونه‌های مستقل،
        تکمیل Metadata، ایجاد مجموعه Validation مستقل و گسترش
        نمونه‌های اصیل، تقلبی و مخلوط‌شده است.
        """
    )


# ============================================================
# TAB 14 — REGION / YEAR SPECIAL ANALYSIS
# ============================================================

with tab_region_year:
    st.subheader("🌍 تحلیل تخصصی منطقه و سال")
    st.caption(
        "این بخش یک تحلیل تکمیلی اکتشافی است و نتیجه اصلی مدل ۱۱ منطقه‌ای را تغییر نمی‌دهد. مدل‌های منطقه‌ای با ۴۰۱ ویژگی مهندسی‌شده اجرا شده‌اند؛ تحلیل شباهت بین سال‌ها بر پایه فضای PCA طیف خام انجام شده است."
    )

    summary = read_csv(str(REGION_YEAR_SUMMARY_PATH))
    region_distribution = read_csv(str(REGION_DISTRIBUTION_PATH))
    regional_results = read_csv(str(REGIONAL_MODEL_RESULTS_PATH))
    similarity = read_csv(str(REGION_YEAR_SIMILARITY_PATH))
    pairwise = read_csv(str(REGION_YEAR_PAIRWISE_PATH))

    region_ready = any(
        not df.empty for df in [summary, region_distribution, regional_results, similarity]
    )

    if not region_ready:
        st.warning(
            """
            تحلیل تخصصی منطقه و سال هنوز اجرا نشده است.

            در ریشه پروژه اجرا کنید:

            `python region_year_analysis.py`
            """
        )
    else:
        st.markdown("### وضعیت اجرای تحلیل")
        s = summary.iloc[0] if not summary.empty else pd.Series(dtype=object)
        independent_samples = (
            s.get("IndependentSamples")
            if "IndependentSamples" in summary.columns
            else dataset["independent_samples"]
        )
        selected_regions = (
            s.get("SelectedRegions")
            if "SelectedRegions" in summary.columns
            else len(region_distribution)
            if not region_distribution.empty
            else "—"
        )
        common_regions = (
            s.get("CommonRegionsBothYears")
            if "CommonRegionsBothYears" in summary.columns
            else 7
        )
        feature_count = (
            s.get("FeatureCount")
            if "FeatureCount" in summary.columns
            else 401
        )
        feature_source = (
            s.get("FeatureSource")
            if "FeatureSource" in summary.columns
            else "engineered_features.csv"
        )

        c1, c2, c3, c4, c5 = st.columns(5)
        with c1:
            st.metric("نمونه‌های مستقل", str(int(float(independent_samples))) if pd.notna(independent_samples) else "—")
        with c2:
            st.metric("مناطق منتخب", str(int(float(selected_regions))) if pd.notna(selected_regions) and str(selected_regions) != "—" else "—")
        with c3:
            st.metric("مناطق دارای هر دو سال", str(int(float(common_regions))) if pd.notna(common_regions) else "—")
        with c4:
            st.metric("ویژگی‌های مدل منطقه‌ای", str(int(float(feature_count))) if pd.notna(feature_count) else "401")
        with c5:
            st.metric("تعداد نقاط طیفی", str(dataset["spectral_points"]))

        st.caption(f"منبع ویژگی‌های مدل منطقه‌ای: **{feature_source}**")
        st.divider()
        st.markdown("### ۱. مدل اختصاصی مناطق دارای نمونه بیشتر")
        st.write(
            "برای کاهش اثر عدم‌توازن شدید در مدل ۱۱ منطقه‌ای، سه منطقه دارای بیشترین تعداد نمونه به‌صورت جداگانه در برابر سایر مناطق بررسی شده‌اند. این مدل‌ها صرفاً شواهد اکتشافی ارائه می‌کنند."
        )

        if not region_distribution.empty:
            st.markdown("#### مناطق منتخب")
            distribution_display = region_distribution.copy()
            rename_distribution = {
                "group": "کد منطقه", "Group": "کد منطقه", "Region": "کد منطقه",
                "RegionLabel": "منطقه", "Samples": "تعداد نمونه", "N": "تعداد نمونه",
            }
            distribution_display = distribution_display.rename(
                columns={k: v for k, v in rename_distribution.items() if k in distribution_display.columns}
            )
            st.dataframe(distribution_display, hide_index=True, width="stretch")

        if regional_results.empty:
            st.info("نتایج مدل‌های اختصاصی مناطق موجود نیست.")
        else:
            rr = regional_results.copy()
            if "Region" not in rr.columns:
                for candidate in ["group", "Group", "region"]:
                    if candidate in rr.columns:
                        rr["Region"] = rr[candidate].astype(str)
                        break
            if "RegionLabel" not in rr.columns:
                rr["RegionLabel"] = rr["Region"] if "Region" in rr.columns else "—"
            for col in [
                "Samples", "PositiveSamples", "TopK", "AccuracyMean", "AccuracyStd",
                "BalancedAccuracyMean", "BalancedAccuracyStd", "MacroF1Mean", "MacroF1Std",
            ]:
                if col in rr.columns:
                    rr[col] = pd.to_numeric(rr[col], errors="coerce")

            if "Region" in rr.columns and "BalancedAccuracyMean" in rr.columns:
                best_per_region = (
                    rr.sort_values(["Region", "BalancedAccuracyMean"], ascending=[True, False])
                    .groupby("Region", as_index=False)
                    .head(1)
                    .sort_values("BalancedAccuracyMean", ascending=False)
                    .reset_index(drop=True)
                )

                st.markdown("#### بهترین مدل هر منطقه")
                best_display_columns = [
                    "Region", "RegionLabel", "Samples", "PositiveSamples", "Model",
                    "TopK", "AccuracyMean", "BalancedAccuracyMean", "MacroF1Mean",
                ]
                best_display_columns = [c for c in best_display_columns if c in best_per_region.columns]
                best_display = best_per_region[best_display_columns].copy()
                rename_best = {
                    "Region": "کد منطقه", "RegionLabel": "منطقه", "Samples": "کل نمونه",
                    "PositiveSamples": "نمونه مثبت", "Model": "مدل", "TopK": "Top-K",
                    "AccuracyMean": "Accuracy", "BalancedAccuracyMean": "Balanced Accuracy",
                    "MacroF1Mean": "Macro F1",
                }
                best_display = best_display.rename(
                    columns={k: v for k, v in rename_best.items() if k in best_display.columns}
                )
                st.dataframe(best_display, hide_index=True, width="stretch")

                k1, k2, k3 = st.columns(3)
                top_region = best_per_region.iloc[0]
                with k1:
                    st.metric("قوی‌ترین سیگنال", str(top_region.get("RegionLabel", top_region.get("Region", "—"))))
                with k2:
                    st.metric("بیشترین Balanced Accuracy", fmt(top_region["BalancedAccuracyMean"]))
                with k3:
                    st.metric(
                        "مدل برتر",
                        f"{top_region.get('Model', '—')} / Top-{int(top_region['TopK'])}"
                        if pd.notna(top_region.get("TopK", float("nan")))
                        else str(top_region.get("Model", "—")),
                    )

                st.markdown("#### Balanced Accuracy مناطق منتخب")
                chart = (
                    best_per_region[["RegionLabel", "BalancedAccuracyMean"]]
                    .dropna()
                    .set_index("RegionLabel")
                )
                if not chart.empty:
                    st.bar_chart(chart)

                if {"RegionLabel", "BalancedAccuracyMean", "BalancedAccuracyStd"}.issubset(best_per_region.columns):
                    st.markdown("#### Balanced Accuracy به همراه عدم‌قطعیت")
                    uncertainty_display = (
                        best_per_region[
                            ["RegionLabel", "BalancedAccuracyMean", "BalancedAccuracyStd"]
                        ]
                        .copy()
                        .rename(
                            columns={
                                "RegionLabel": "منطقه",
                                "BalancedAccuracyMean": "میانگین BA",
                                "BalancedAccuracyStd": "انحراف معیار BA",
                            }
                        )
                    )
                    st.dataframe(uncertainty_display.round(4), hide_index=True, width="stretch")

            st.markdown("#### تمام نتایج مدل‌های منطقه‌ای")
            full_display = rr.copy()
            display_cols = [
                c for c in [
                    "Region", "RegionLabel", "Samples", "PositiveSamples", "Model", "TopK",
                    "AccuracyMean", "AccuracyStd", "BalancedAccuracyMean", "BalancedAccuracyStd",
                    "MacroF1Mean", "MacroF1Std",
                ] if c in full_display.columns
            ]
            if display_cols:
                st.dataframe(full_display[display_cols].round(4), hide_index=True, width="stretch")
            st.info(
                "مدل‌های منطقه‌ای در این داده‌ها فقط یک سیگنال اکتشافی از قابلیت تفکیک هر منطقه در برابر سایر مناطق نشان می‌دهند. به‌دلیل تعداد کم نمونه‌ها و عدم‌توازن بین مناطق، این اعداد مدل عملیاتی یا اثبات قطعی منشأ جغرافیایی نیستند."
            )

        st.divider()
        st.markdown("### ۲. بررسی پایداری الگوی منطقه‌ای بین ۱۳۹۴ و ۱۴۰۴")
        st.write(
            "در این آزمون بررسی می‌شود آیا نمونه‌های یک منطقه در سال ۱۴۰۴، در فضای طیفی به نمونه‌های همان منطقه در سال ۱۳۹۴ نزدیک‌ترند یا خیر. این بخش از تحلیل با PCA تازه و مستقل از مدل‌های منطقه‌ای محاسبه شده است."
        )

        if similarity.empty:
            st.info("نتایج Region × Year موجود نیست.")
        else:
            sim = similarity.copy()
            aliases = {
                "SameRegionDistanceMean": ["SameRegionDistanceMean", "MeanSameRegionDistance"],
                "OtherRegionDistanceMean": ["OtherRegionDistanceMean", "MeanOtherRegionDistance"],
                "SameToOtherDistanceRatio": ["SameToOtherDistanceRatio"],
                "NearestSameRegionFraction": ["NearestSameRegionFraction"],
                "PermutationPValue": ["PermutationPValue"],
            }
            for canonical, candidates in aliases.items():
                if canonical not in sim.columns:
                    found = next((c for c in candidates if c in sim.columns), None)
                    if found is not None:
                        sim[canonical] = sim[found]
            if "Region" not in sim.columns:
                for candidate in ["group", "Group"]:
                    if candidate in sim.columns:
                        sim["Region"] = sim[candidate].astype(str)
                        break
            if "RegionLabel" not in sim.columns:
                sim["RegionLabel"] = sim["Region"] if "Region" in sim.columns else "—"
            for col in [
                "SameRegionDistanceMean", "OtherRegionDistanceMean",
                "SameToOtherDistanceRatio", "NearestSameRegionFraction", "PermutationPValue",
            ]:
                if col in sim.columns:
                    sim[col] = pd.to_numeric(sim[col], errors="coerce")

            overall = None
            if "Region" in sim.columns:
                overall_rows = sim[sim["Region"].astype(str).str.upper() == "ALL"]
                if not overall_rows.empty:
                    overall = overall_rows.iloc[0]

            if overall is not None:
                same_mean = overall.get("SameRegionDistanceMean", float("nan"))
                other_mean = overall.get("OtherRegionDistanceMean", float("nan"))
                ratio = overall.get("SameToOtherDistanceRatio", float("nan"))
                nearest_fraction = overall.get("NearestSameRegionFraction", float("nan"))
                p_value = overall.get("PermutationPValue", float("nan"))
                c1, c2, c3, c4, c5 = st.columns(5)
                with c1:
                    st.metric("میانگین فاصله همان منطقه", fmt(same_mean, 2))
                with c2:
                    st.metric("میانگین فاصله سایر مناطق", fmt(other_mean, 2))
                with c3:
                    st.metric("نسبت Same / Other", fmt(ratio, 4))
                with c4:
                    st.metric(
                        "Nearest Same-Region",
                        f"{fmt(nearest_fraction * 100, 1)}%" if pd.notna(nearest_fraction) else "—",
                    )
                with c5:
                    st.metric("Permutation p-value", fmt(p_value, 4))

                if pd.notna(ratio) and pd.notna(nearest_fraction) and pd.notna(p_value):
                    if ratio < 1 and nearest_fraction > 0.5 and p_value < 0.05:
                        st.success("الگوی کلی تحلیل با پایداری منطقه‌ای بین دو سال سازگار است.")
                    else:
                        st.warning("شواهد کافی برای پایداری یک الگوی منطقه‌ای بین سال‌های ۱۳۹۴ و ۱۴۰۴ مشاهده نشد.")

            st.divider()
            st.markdown("#### نتایج منطقه‌ای شباهت بین دو سال")
            sim_display = sim.copy().rename(
                columns={
                    "Region": "کد منطقه", "RegionLabel": "منطقه", "N_1394": "نمونه ۱۳۹۴", "N_1404": "نمونه ۱۴۰۴",
                    "SameRegionDistanceMean": "فاصله همان منطقه", "OtherRegionDistanceMean": "فاصله سایر مناطق",
                    "SameToOtherDistanceRatio": "نسبت Same/Other", "NearestSameRegionFraction": "نزدیک‌ترین هم‌منطقه",
                    "PermutationPValue": "p-value", "EvidenceRatioBelow1": "Ratio < 1", "EvidenceNearestAbove05": "Nearest > 0.5",
                }
            )
            for col in ["فاصله همان منطقه", "فاصله سایر مناطق", "نسبت Same/Other", "نزدیک‌ترین هم‌منطقه", "p-value"]:
                if col in sim_display.columns:
                    sim_display[col] = pd.to_numeric(sim_display[col], errors="coerce").round(4)
            st.dataframe(sim_display, hide_index=True, width="stretch")

            if {"RegionLabel", "SameRegionDistanceMean", "OtherRegionDistanceMean"}.issubset(sim.columns):
                st.markdown("#### فاصله نمونه‌های همان منطقه در برابر سایر مناطق")
                region_distance_chart = (
                    sim[sim["Region"].astype(str).str.upper() != "ALL"]
                    [["RegionLabel", "SameRegionDistanceMean", "OtherRegionDistanceMean"]]
                    .set_index("RegionLabel")
                )
                if not region_distance_chart.empty:
                    st.bar_chart(region_distance_chart)

            if {"RegionLabel", "NearestSameRegionFraction"}.issubset(sim.columns):
                st.markdown("#### سهم نزدیک‌ترین نمونه از همان منطقه")
                nearest_chart = (
                    sim[sim["Region"].astype(str).str.upper() != "ALL"]
                    [["RegionLabel", "NearestSameRegionFraction"]]
                    .set_index("RegionLabel")
                )
                if not nearest_chart.empty:
                    st.bar_chart(nearest_chart)

            if overall is not None:
                ratio = overall.get("SameToOtherDistanceRatio", float("nan"))
                nearest_fraction = overall.get("NearestSameRegionFraction", float("nan"))
                p_value = overall.get("PermutationPValue", float("nan"))
                st.divider()
                st.markdown("#### جمع‌بندی خودکار این آزمون")
                if pd.notna(ratio) and pd.notna(nearest_fraction) and pd.notna(p_value):
                    if ratio < 1 and nearest_fraction > 0.5 and p_value < 0.05:
                        st.success("این سه معیار به‌طور هم‌زمان از پایداری الگوی منطقه‌ای بین دو سال پشتیبانی می‌کنند.")
                    else:
                        st.warning("در اجرای فعلی، شواهد کافی برای تأیید پایداری یک امضای منطقه‌ای بین دو سال مشاهده نشد.")
                    st.markdown(
                        f"**نسبت فاصله:** `{fmt(ratio, 4)}`\n\n**سهم نزدیک‌ترین هم‌منطقه:** `{fmt(nearest_fraction * 100, 1)}%`\n\n**Permutation p-value:** `{fmt(p_value, 4)}`"
                    )
            st.info(
                "تفسیر علمی: نسبت کمتر از ۱ به این معناست که فاصله نمونه‌های همان منطقه از نمونه‌های سال دیگر کمتر از فاصله نمونه‌های مناطق دیگر بوده است."
            )

        st.divider()
        st.markdown("### ۳. جمع‌بندی قابل ارائه در گزارش")
        st.warning(
            """
            **نتیجه پیشنهادی:**

            مدل‌های منطقه‌ای اختصاصی با ۴۰۱ ویژگی مهندسی‌شده نشان دادند که در میان سه منطقه دارای بیشترین نمونه، اراک قوی‌ترین سیگنال اکتشافی را دارد؛ با این حال محدودیت تعداد نمونه‌ها اجازه نمی‌دهد این سیگنال به‌عنوان مدل عملیاتی یا اثبات قطعی منشأ معرفی شود.

            از طرف دیگر، آزمون شباهت بین سال‌های ۱۳۹۴ و ۱۴۰۴ شواهد کافی برای وجود یک امضای طیفی پایدار و مستقل از سال در مناطق مشترک ارائه نکرد. بنابراین در شرایط فعلی، نمی‌توان گفت که نمونه‌های یک شهر در دو سال مختلف الزاماً به یکدیگر نزدیک‌تر از نمونه‌های سایر شهرها هستند.
            """
        )
        if not pairwise.empty:
            with st.expander("▶ مشاهده تمام فاصله‌های زوجی بین ۱۳۹۴ و ۱۴۰۴"):
                st.dataframe(pairwise, hide_index=True, width="stretch")


# ============================================================
# REFERENCE DOCUMENT HELPERS
# ============================================================

@st.cache_data(show_spinner=False)
def read_reference_docx(path_string: str):
    path = Path(path_string)
    if not path.exists():
        return None, []
    try:
        from docx import Document
        document = Document(str(path))
        blocks = []
        for paragraph in document.paragraphs:
            value = paragraph.text.strip()
            if not value:
                continue
            style_name = ""
            try:
                style_name = (paragraph.style.name or "").lower()
            except Exception:
                pass
            if "heading 1" in style_name:
                kind = "h1"
            elif "heading 2" in style_name:
                kind = "h2"
            elif "heading 3" in style_name:
                kind = "h3"
            elif "heading" in style_name:
                kind = "h4"
            else:
                kind = "p"
            blocks.append((kind, value))

        tables = []
        for table in document.tables:
            rows = []
            for row in table.rows:
                values = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                if any(values):
                    rows.append(values)
            if rows:
                tables.append(rows)
        return str(path), {"blocks": blocks, "tables": tables}
    except Exception as exc:
        return str(path), {"error": str(exc), "blocks": [], "tables": []}


# ============================================================
# TAB 15 — REFERENCE PROJECT DOCUMENT
# ============================================================

with tab_reference_doc:
    st.subheader("سند مرجع پروژه")
    st.caption(
        "سند Word اصلی پروژه از مسیر docs/saffron-05.docx خوانده می‌شود و محتوای آن مستقیماً در داشبورد نمایش داده می‌شود."
    )
    ref_path = REFERENCE_DOCX_PATH
    ref_path_string, reference_data = read_reference_docx(str(ref_path))

    if ref_path.exists():
        st.success("✅ فایل saffron-05.docx با موفقیت از پوشه docs خوانده شد.")
        c1, c2 = st.columns([1, 3])
        with c1:
            st.metric("وضعیت سند", "در دسترس")
        with c2:
            st.write(f"مسیر: **{ref_path}**")

        st.download_button(
            label="⬇️ دانلود سند اصلی Word",
            data=ref_path.read_bytes(),
            file_name=ref_path.name,
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            width="content",
        )

        if reference_data and isinstance(reference_data, dict):
            if reference_data.get("error"):
                st.error("فایل وجود دارد، اما در خواندن محتوای Word خطا رخ داده است.")
                st.code(reference_data["error"], language="text")
            else:
                st.divider()
                st.markdown("### محتوای سند مرجع")
                st.info(
                    "این بخش همان سند طراحی علمی و فنی اولیه پروژه است. نتایج واقعی اجرای فعلی باید از تب‌های تحلیلی و فایل‌های خروجی پایپ‌لاین خوانده شوند."
                )
                for kind, value in reference_data["blocks"]:
                    if kind == "h1":
                        st.markdown(f"# {value}")
                    elif kind == "h2":
                        st.markdown(f"## {value}")
                    elif kind == "h3":
                        st.markdown(f"### {value}")
                    elif kind == "h4":
                        st.markdown(f"#### {value}")
                    else:
                        st.markdown(value)

                for table_index, rows in enumerate(reference_data["tables"], start=1):
                    st.markdown(f"### جدول {table_index}")
                    if rows:
                        width = max(len(row) for row in rows)
                        normalized_rows = [row + [""] * (width - len(row)) for row in rows]
                        if len(normalized_rows) >= 2:
                            columns = normalized_rows[0]
                            body = normalized_rows[1:]
                            clean_columns = []
                            seen = {}
                            for index, column in enumerate(columns):
                                name = column.strip() or f"ستون {index + 1}"
                                count = seen.get(name, 0)
                                seen[name] = count + 1
                                if count:
                                    name = f"{name} ({count + 1})"
                                clean_columns.append(name)
                            table_df = pd.DataFrame(body, columns=clean_columns)
                            st.dataframe(table_df, hide_index=True, width="stretch")
                        else:
                            st.dataframe(pd.DataFrame(normalized_rows), hide_index=True, width="stretch")
    else:
        st.warning("⚠️ فایل docs/saffron-05.docx در کنار پروژه پیدا نشد.")
        st.write("برای نمایش خودکار متن، فایل باید دقیقاً در این مسیر قرار داشته باشد:")
        st.code(str(ref_path), language="text")
        st.info("در این وضعیت، ابتدا فایل Word را در پوشه docs قرار دهید و سپس داشبورد را Refresh کنید.")
