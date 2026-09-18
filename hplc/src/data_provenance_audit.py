from __future__ import annotations

from pathlib import Path
from collections import Counter

import numpy as np
import pandas as pd
from openpyxl import load_workbook


# ==============================================================================
# PATHS
# ==============================================================================

ROOT = Path(__file__).resolve().parents[1]

EXCEL_PATH = ROOT / "data" / "raw" / "hplc.xlsx"
MAPPING_PATH = ROOT / "data" / "mapping" / "sample_mapping.csv"
OUTPUT_DIR = ROOT / "outputs"

OUTPUT_DIR.mkdir(parents=True, exist_ok=True)


# ==============================================================================
# SETTINGS
# ==============================================================================

SHEETS = {
    "Crocin 440": 440,
    "picrocrocin 250": 250,
    "safranal 308": 308,
}

MAX_TIME = 30.0

DUPLICATE_PAIRS = [
    (1, 2),
    (24, 25),
    (28, 29),
    (34, 35),
]


# ==============================================================================
# HELPERS
# ==============================================================================

def normalize_text(value) -> str:
    if value is None:
        return ""

    return str(value).strip()


def safe_float(value):
    try:
        result = float(value)

        if np.isfinite(result):
            return result

    except Exception:
        pass

    return None


def find_header_row(ws) -> int:
    """
    Find the row containing Number + Group.
    openpyxl rows are 1-based.
    """

    for row_idx in range(
        1,
        min(ws.max_row, 30) + 1,
    ):

        values = [
            normalize_text(
                ws.cell(row_idx, col_idx).value
            ).lower()
            for col_idx in range(
                1,
                min(ws.max_column, 20) + 1,
            )
        ]

        has_number = any(
            value in {
                "number",
                "sample number",
            }
            for value in values
        )

        has_group = "group" in values

        if has_number and has_group:
            return row_idx

    raise ValueError(
        f"Header row not found in sheet '{ws.title}'"
    )


def detect_number_column(
    ws,
    header_row: int,
) -> int:

    for col_idx in range(
        1,
        ws.max_column + 1,
    ):

        text = normalize_text(
            ws.cell(
                header_row,
                col_idx,
            ).value
        ).lower()

        if text in {
            "number",
            "sample number",
        }:
            return col_idx

    # Known fallback from project workbook structure
    return 3


def detect_group_column(
    ws,
    header_row: int,
) -> int:

    for col_idx in range(
        1,
        ws.max_column + 1,
    ):

        text = normalize_text(
            ws.cell(
                header_row,
                col_idx,
            ).value
        ).lower()

        if text == "group":
            return col_idx

    # Known fallback from project workbook structure
    return 4


def detect_time_columns(
    ws,
    header_row: int,
) -> tuple[list[int], list[float]]:

    columns = []
    times = []

    for col_idx in range(
        1,
        ws.max_column + 1,
    ):

        value = ws.cell(
            header_row,
            col_idx,
        ).value

        numeric = safe_float(
            value
        )

        if numeric is None:
            continue

        if (
            numeric > 0.0
            and numeric <= MAX_TIME
        ):
            columns.append(
                col_idx
            )

            times.append(
                numeric
            )

    if not columns:
        raise ValueError(
            f"No time columns found in "
            f"sheet '{ws.title}'"
        )

    order = np.argsort(
        np.asarray(times)
    )

    ordered_columns = [
        columns[i]
        for i in order
    ]

    ordered_times = [
        times[i]
        for i in order
    ]

    return (
        ordered_columns,
        ordered_times,
    )


def read_mapping() -> pd.DataFrame:

    if not MAPPING_PATH.exists():
        raise FileNotFoundError(
            f"Mapping file not found:\n{MAPPING_PATH}"
        )

    mapping = pd.read_csv(
        MAPPING_PATH
    )

    normalized = {
        str(col)
        .strip()
        .lower()
        .replace("_", "")
        .replace(" ", ""):
            col
        for col in mapping.columns
    }

    sample_col = None
    group_col = None
    year_col = None

    for normalized_name, real_name in normalized.items():

        if normalized_name in {
            "sampleid",
            "sample",
            "number",
        }:
            sample_col = real_name

        if normalized_name in {
            "group",
            "class",
            "region",
        }:
            group_col = real_name

        if normalized_name in {
            "year",
            "harvestyear",
            "harvest",
        }:
            year_col = real_name

    if sample_col is None:
        raise ValueError(
            "Sample ID column not found in mapping CSV."
        )

    if group_col is None:
        raise ValueError(
            "Group column not found in mapping CSV."
        )

    result = pd.DataFrame()

    result["SampleId"] = pd.to_numeric(
        mapping[sample_col],
        errors="coerce",
    )

    result["ActualClass"] = (
        mapping[group_col]
        .astype(str)
        .str.strip()
    )

    if year_col is not None:
        result["Year"] = mapping[
            year_col
        ]
    else:
        result["Year"] = ""

    result = result.dropna(
        subset=["SampleId"]
    )

    result["SampleId"] = (
        result["SampleId"]
        .astype(int)
    )

    return result


# ==============================================================================
# SAMPLE SIGNAL EXTRACTION
# ==============================================================================

def extract_sheet(
    ws,
    wavelength: int,
    mapping: pd.DataFrame,
):

    header_row = find_header_row(
        ws
    )

    number_col = detect_number_column(
        ws,
        header_row,
    )

    group_col = detect_group_column(
        ws,
        header_row,
    )

    time_columns, times = detect_time_columns(
        ws,
        header_row,
    )

    mapping_index = (
        mapping
        .drop_duplicates(
            subset=["SampleId"]
        )
        .set_index("SampleId")
    )

    records = []

    for row_idx in range(
        header_row + 1,
        ws.max_row + 1,
    ):

        sample_value = ws.cell(
            row_idx,
            number_col,
        ).value

        sample_numeric = safe_float(
            sample_value
        )

        if sample_numeric is None:
            continue

        sample_id = int(
            sample_numeric
        )

        group_value = normalize_text(
            ws.cell(
                row_idx,
                group_col,
            ).value
        )

        values = []

        formula_count = 0
        blank_count = 0
        text_count = 0
        error_count = 0
        numeric_count = 0

        number_formats = Counter()
        style_ids = Counter()
        data_types = Counter()

        for col_idx in time_columns:

            cell = ws.cell(
                row_idx,
                col_idx,
            )

            value = cell.value

            data_types[
                cell.data_type
            ] += 1

            number_formats[
                cell.number_format
            ] += 1

            style_ids[
                cell.style_id
            ] += 1

            if cell.data_type == "f":
                formula_count += 1

            if value is None:
                blank_count += 1
                continue

            if (
                isinstance(value, str)
                and value.startswith("=")
            ):
                formula_count += 1

            numeric = safe_float(
                value
            )

            if numeric is not None:

                numeric_count += 1
                values.append(
                    numeric
                )

            elif isinstance(value, str):

                text_count += 1

                if value.startswith(
                    "#"
                ):
                    error_count += 1

        if len(values) == 0:
            continue

        y = np.asarray(
            values,
            dtype=float,
        )

        current_times = np.asarray(
            times[:len(y)],
            dtype=float,
        )

        abs_y = np.abs(y)

        auc = float(
            np.trapezoid(
                y,
                current_times,
            )
        )

        abs_auc = float(
            np.trapezoid(
                abs_y,
                current_times,
            )
        )

        median_abs = float(
            np.median(
                abs_y
            )
        )

        mean_abs = float(
            np.mean(
                abs_y
            )
        )

        max_abs = float(
            np.max(
                abs_y
            )
        )

        std_abs = float(
            np.std(
                abs_y
            )
        )

        min_value = float(
            np.min(y)
        )

        max_value = float(
            np.max(y)
        )

        # --------------------------------------------------------------
        # Mapping
        # --------------------------------------------------------------

        mapped_group = ""

        mapped_year = ""

        if sample_id in mapping_index.index:

            mapped_group = normalize_text(
                mapping_index.loc[
                    sample_id,
                    "ActualClass",
                ]
            )

            mapped_year = str(
                mapping_index.loc[
                    sample_id,
                    "Year",
                ]
            )

        records.append(
            {
                "SampleId": sample_id,
                "Wavelength": wavelength,
                "WorkbookGroup": group_value,
                "MappedGroup": mapped_group,
                "Year": mapped_year,
                "AUC": auc,
                "AbsAUC": abs_auc,
                "MedianAbs": median_abs,
                "MeanAbs": mean_abs,
                "MaxAbs": max_abs,
                "StdAbs": std_abs,
                "MinIntensity": min_value,
                "MaxIntensity": max_value,
                "SignalPoints": len(y),
                "BlankCells": blank_count,
                "NumericCells": numeric_count,
                "TextCells": text_count,
                "ErrorCells": error_count,
                "FormulaCells": formula_count,
                "NumberFormatCount": len(
                    number_formats
                ),
                "StyleIdCount": len(
                    style_ids
                ),
                "DataTypePattern": (
                    "|".join(
                        f"{k}:{v}"
                        for k, v
                        in sorted(
                            data_types.items()
                        )
                    )
                ),
            }
        )

    return (
        pd.DataFrame(records),
        {
            "sheet": ws.title,
            "wavelength": wavelength,
            "header_row": header_row,
            "number_col": number_col,
            "group_col": group_col,
            "time_columns": len(
                time_columns
            ),
            "time_min": float(
                min(times)
            ),
            "time_max": float(
                max(times)
            ),
            "time_median_step": float(
                np.median(
                    np.diff(times)
                )
            )
            if len(times) > 1
            else np.nan,
            "time_min_step": float(
                np.min(
                    np.diff(times)
                )
            )
            if len(times) > 1
            else np.nan,
            "time_max_step": float(
                np.max(
                    np.diff(times)
                )
            )
            if len(times) > 1
            else np.nan,
        },
    )


# ==============================================================================
# YEAR SUMMARY
# ==============================================================================

def year_summary(
    sample_df: pd.DataFrame,
) -> pd.DataFrame:

    numeric = [
        "AUC",
        "AbsAUC",
        "MedianAbs",
        "MeanAbs",
        "MaxAbs",
        "StdAbs",
        "MinIntensity",
        "MaxIntensity",
    ]

    summary = (
        sample_df
        .groupby(
            [
                "Wavelength",
                "Year",
            ]
        )[numeric]
        .agg(
            [
                "median",
                "mean",
                "std",
                "min",
                "max",
            ]
        )
        .reset_index()
    )

    summary.columns = [
        "_".join(
            str(part)
            for part in column
            if str(part) != ""
        )
        for column in summary.columns
    ]

    return summary


# ==============================================================================
# YEAR SCALE RATIOS
# ==============================================================================

def year_scale_ratios(
    sample_df: pd.DataFrame,
) -> pd.DataFrame:

    rows = []

    metrics = [
        "AbsAUC",
        "MedianAbs",
        "MeanAbs",
        "MaxAbs",
    ]

    for wavelength in sorted(
        sample_df[
            "Wavelength"
        ].unique()
    ):

        current = sample_df[
            sample_df[
                "Wavelength"
            ] == wavelength
        ]

        medians = (
            current
            .groupby("Year")[
                metrics
            ]
            .median()
        )

        years = sorted(
            medians.index.astype(str)
        )

        if len(years) != 2:
            continue

        year_a, year_b = years

        for metric in metrics:

            a = float(
                medians.loc[
                    year_a,
                    metric,
                ]
            )

            b = float(
                medians.loc[
                    year_b,
                    metric,
                ]
            )

            if abs(a) > 1e-15:
                ratio = b / a
            else:
                ratio = np.nan

            log10_ratio = np.nan

            if (
                a > 0
                and b > 0
            ):
                log10_ratio = float(
                    np.log10(
                        b / a
                    )
                )

            rows.append(
                {
                    "Wavelength": wavelength,
                    "Metric": metric,
                    "YearA": year_a,
                    "YearB": year_b,
                    "MedianYearA": a,
                    "MedianYearB": b,
                    "YearB_over_YearA": ratio,
                    "Log10_YearB_over_YearA": log10_ratio,
                }
            )

    return pd.DataFrame(
        rows
    )


# ==============================================================================
# SAMPLE SCALE PROFILE
# ==============================================================================

def sample_scale_profile(
    sample_df: pd.DataFrame,
) -> pd.DataFrame:

    pivot = (
        sample_df
        .pivot_table(
            index=[
                "SampleId",
                "MappedGroup",
                "Year",
            ],
            columns="Wavelength",
            values=[
                "MedianAbs",
                "MeanAbs",
                "MaxAbs",
                "AbsAUC",
            ],
        )
    )

    pivot.columns = [
        f"{metric}_{int(wavelength)}"
        for metric, wavelength
        in pivot.columns
    ]

    pivot = (
        pivot
        .reset_index()
    )

    return pivot


# ==============================================================================
# PER-SAMPLE YEAR SCALE INDEX
# ==============================================================================

def calculate_year_normalized_scale(
    profile_df: pd.DataFrame,
) -> pd.DataFrame:

    result = profile_df.copy()

    for wavelength in [250, 308, 440]:

        column = (
            f"MedianAbs_{wavelength}"
        )

        if column not in result.columns:
            continue

        year_medians = (
            result
            .groupby("Year")[
                column
            ]
            .median()
        )

        normalized_column = (
            f"MedianAbs_vs_YearMedian_{wavelength}"
        )

        result[
            normalized_column
        ] = np.nan

        for year in year_medians.index:

            mask = (
                result["Year"]
                .astype(str)
                == str(year)
            )

            denominator = (
                year_medians.loc[year]
            )

            if (
                pd.notna(denominator)
                and abs(
                    float(denominator)
                ) > 1e-15
            ):

                result.loc[
                    mask,
                    normalized_column,
                ] = (
                    result.loc[
                        mask,
                        column,
                    ]
                    / float(denominator)
                )

    return result


# ==============================================================================
# WORKBOOK STRUCTURE
# ==============================================================================

def workbook_structure(
    wb_formula,
    wb_values,
) -> pd.DataFrame:

    rows = []

    for sheet_name, wavelength in SHEETS.items():

        ws_formula = wb_formula[
            sheet_name
        ]

        ws_values = wb_values[
            sheet_name
        ]

        hidden_rows = [
            row_idx
            for row_idx, dimension
            in ws_formula.row_dimensions.items()
            if dimension.hidden
        ]

        hidden_columns = [
            column_letter
            for column_letter, dimension
            in ws_formula.column_dimensions.items()
            if dimension.hidden
        ]

        merged_ranges = [
            str(rng)
            for rng in ws_formula.merged_cells.ranges
        ]

        header_row = find_header_row(
            ws_formula
        )

        number_col = detect_number_column(
            ws_formula,
            header_row,
        )

        group_col = detect_group_column(
            ws_formula,
            header_row,
        )

        time_cols, times = detect_time_columns(
            ws_formula,
            header_row,
        )

        formula_cells = 0
        formula_cached_values = 0
        differing_formula_value_cells = 0

        for row_idx in range(
            header_row + 1,
            ws_formula.max_row + 1,
        ):

            for col_idx in time_cols:

                cell_formula = ws_formula.cell(
                    row_idx,
                    col_idx,
                )

                cell_value = ws_values.cell(
                    row_idx,
                    col_idx,
                )

                if (
                    cell_formula.data_type
                    == "f"
                ):
                    formula_cells += 1

                    if (
                        cell_value.value
                        is not None
                    ):
                        formula_cached_values += 1

                    formula_text = (
                        cell_formula.value
                    )

                    cached = (
                        cell_value.value
                    )

                    if (
                        formula_text
                        != cached
                    ):
                        differing_formula_value_cells += 1

        rows.append(
            {
                "Sheet": sheet_name,
                "Wavelength": wavelength,
                "MaxRow": ws_formula.max_row,
                "MaxColumn": ws_formula.max_column,
                "HeaderRow": header_row,
                "NumberColumn": number_col,
                "GroupColumn": group_col,
                "TimeColumns": len(time_cols),
                "TimeMin": min(times),
                "TimeMax": max(times),
                "MedianTimeStep": (
                    np.median(
                        np.diff(times)
                    )
                    if len(times) > 1
                    else np.nan
                ),
                "MergedRangeCount": len(
                    merged_ranges
                ),
                "HiddenRowCount": len(
                    hidden_rows
                ),
                "HiddenColumnCount": len(
                    hidden_columns
                ),
                "FormulaCellsInSignalRegion": (
                    formula_cells
                ),
                "FormulaCachedValueCells": (
                    formula_cached_values
                ),
                "FormulaValueDifferenceCells": (
                    differing_formula_value_cells
                ),
                "FreezePanes": str(
                    ws_formula.freeze_panes
                ),
                "AutoFilter": (
                    str(
                        ws_formula.auto_filter.ref
                    )
                    if ws_formula.auto_filter.ref
                    else ""
                ),
            }
        )

    return pd.DataFrame(
        rows
    )


# ==============================================================================
# SAMPLE / SHEET CONSISTENCY
# ==============================================================================

def sheet_consistency(
    sample_df: pd.DataFrame,
    mapping: pd.DataFrame,
) -> pd.DataFrame:

    rows = []

    mapping_ids = set(
        mapping["SampleId"]
        .astype(int)
    )

    for wavelength in sorted(
        sample_df[
            "Wavelength"
        ].unique()
    ):

        current = sample_df[
            sample_df[
                "Wavelength"
            ] == wavelength
        ]

        sheet_ids = set(
            current["SampleId"]
            .astype(int)
        )

        missing_in_sheet = sorted(
            mapping_ids - sheet_ids
        )

        extra_in_sheet = sorted(
            sheet_ids - mapping_ids
        )

        duplicate_ids = sorted(
            [
                int(sample_id)
                for sample_id, count
                in current[
                    "SampleId"
                ]
                .value_counts()
                .items()
                if count > 1
            ]
        )

        workbook_group_mismatch = (
            current[
                current[
                    "WorkbookGroup"
                ].astype(str)
                != current[
                    "MappedGroup"
                ].astype(str)
            ]
        )

        rows.append(
            {
                "Wavelength": wavelength,
                "SheetSampleCount": len(
                    sheet_ids
                ),
                "MappingSampleCount": len(
                    mapping_ids
                ),
                "MissingMappingSamples": (
                    ",".join(
                        map(
                            str,
                            missing_in_sheet,
                        )
                    )
                ),
                "ExtraWorkbookSamples": (
                    ",".join(
                        map(
                            str,
                            extra_in_sheet,
                        )
                    )
                ),
                "DuplicateSampleIdsInSheet": (
                    ",".join(
                        map(
                            str,
                            duplicate_ids,
                        )
                    )
                ),
                "WorkbookVsMappingGroupMismatchCount": (
                    len(
                        workbook_group_mismatch
                    )
                ),
                "WorkbookVsMappingMismatchIds": (
                    ",".join(
                        map(
                            str,
                            sorted(
                                workbook_group_mismatch[
                                    "SampleId"
                                ].tolist()
                            ),
                        )
                    )
                ),
            }
        )

    return pd.DataFrame(
        rows
    )


# ==============================================================================
# CELL FORMAT CONSISTENCY
# ==============================================================================

def format_consistency(
    wb,
) -> pd.DataFrame:

    rows = []

    for sheet_name, wavelength in SHEETS.items():

        ws = wb[
            sheet_name
        ]

        header_row = find_header_row(
            ws
        )

        time_cols, _ = detect_time_columns(
            ws,
            header_row,
        )

        number_formats = Counter()
        style_ids = Counter()
        data_types = Counter()

        for row_idx in range(
            header_row + 1,
            ws.max_row + 1,
        ):

            for col_idx in time_cols:

                cell = ws.cell(
                    row_idx,
                    col_idx,
                )

                number_formats[
                    cell.number_format
                ] += 1

                style_ids[
                    cell.style_id
                ] += 1

                data_types[
                    cell.data_type
                ] += 1

        rows.append(
            {
                "Sheet": sheet_name,
                "Wavelength": wavelength,
                "UniqueNumberFormats": len(
                    number_formats
                ),
                "TopNumberFormats": (
                    str(
                        number_formats.most_common(
                            10
                        )
                    )
                ),
                "UniqueStyleIds": len(
                    style_ids
                ),
                "TopStyleIds": (
                    str(
                        style_ids.most_common(
                            10
                        )
                    )
                ),
                "DataTypes": (
                    str(
                        dict(
                            data_types
                        )
                    )
                ),
            }
        )

    return pd.DataFrame(
        rows
    )


# ==============================================================================
# YEAR-SPECIFIC SAMPLE EXAMPLES
# ==============================================================================

def representative_samples(
    sample_df: pd.DataFrame,
) -> pd.DataFrame:

    rows = []

    for wavelength in sorted(
        sample_df[
            "Wavelength"
        ].unique()
    ):

        current = sample_df[
            sample_df[
                "Wavelength"
            ] == wavelength
        ].copy()

        for year in sorted(
            current[
                "Year"
            ].astype(str).unique()
        ):

            subset = current[
                current[
                    "Year"
                ].astype(str)
                == str(year)
            ].copy()

            subset = subset.sort_values(
                "MedianAbs"
            )

            if subset.empty:
                continue

            first = subset.iloc[0]
            middle = subset.iloc[
                len(subset) // 2
            ]
            last = subset.iloc[-1]

            for label, row in [
                ("Low", first),
                ("Median", middle),
                ("High", last),
            ]:

                rows.append(
                    {
                        "Wavelength": wavelength,
                        "Year": year,
                        "Position": label,
                        "SampleId": int(
                            row["SampleId"]
                        ),
                        "MappedGroup": row[
                            "MappedGroup"
                        ],
                        "MedianAbs": row[
                            "MedianAbs"
                        ],
                        "MaxAbs": row[
                            "MaxAbs"
                        ],
                        "AbsAUC": row[
                            "AbsAUC"
                        ],
                        "SignalPoints": row[
                            "SignalPoints"
                        ],
                    }
                )

    return pd.DataFrame(
        rows
    )


# ==============================================================================
# PRINT
# ==============================================================================

def print_section(
    title: str,
):

    print("\n")
    print("=" * 78)
    print(title)
    print("=" * 78)


# ==============================================================================
# MAIN
# ==============================================================================

def main():

    print("=" * 78)
    print(
        "HPLC DATA PROVENANCE / FORENSIC AUDIT"
    )
    print("=" * 78)

    # ------------------------------------------------------------------
    # Existence
    # ------------------------------------------------------------------

    if not EXCEL_PATH.exists():
        raise FileNotFoundError(
            f"HPLC workbook not found:\n{EXCEL_PATH}"
        )

    if not MAPPING_PATH.exists():
        raise FileNotFoundError(
            f"Mapping file not found:\n{MAPPING_PATH}"
        )

    print(
        f"\nWorkbook: {EXCEL_PATH}"
    )

    print(
        f"Mapping : {MAPPING_PATH}"
    )

    # ------------------------------------------------------------------
    # Load workbook twice
    # ------------------------------------------------------------------

    print(
        "\nLoading workbook (formulas)..."
    )

    wb_formula = load_workbook(
        EXCEL_PATH,
        data_only=False,
        read_only=False,
    )

    print(
        "Loading workbook (cached values)..."
    )

    wb_values = load_workbook(
        EXCEL_PATH,
        data_only=True,
        read_only=False,
    )

    mapping = read_mapping()

    print(
        f"\nMapping rows: {len(mapping)}"
    )

    print(
        f"Unique mapping SampleIds: "
        f"{mapping['SampleId'].nunique()}"
    )

    # ------------------------------------------------------------------
    # Sheet names
    # ------------------------------------------------------------------

    print_section(
        "WORKBOOK SHEETS"
    )

    print(
        "Workbook sheets:"
    )

    for name in wb_formula.sheetnames:
        print(
            f"  - {name}"
        )

    missing_sheets = [
        sheet
        for sheet in SHEETS
        if sheet not in wb_formula.sheetnames
    ]

    if missing_sheets:

        print(
            "\nMissing expected sheets:"
        )

        for sheet in missing_sheets:
            print(
                f"  !! {sheet}"
            )

        raise ValueError(
            "Expected HPLC sheets are missing."
        )

    # ------------------------------------------------------------------
    # Workbook structure
    # ------------------------------------------------------------------

    structure_df = workbook_structure(
        wb_formula,
        wb_values,
    )

    print_section(
        "WORKBOOK STRUCTURE"
    )

    print(
        structure_df.to_string(
            index=False
        )
    )

    # ------------------------------------------------------------------
    # Extract signal data
    # ------------------------------------------------------------------

    all_samples = []
    sheet_infos = []

    for sheet_name, wavelength in SHEETS.items():

        print(
            f"\nAuditing sheet: "
            f"{sheet_name}"
        )

        df, info = extract_sheet(
            wb_formula[
                sheet_name
            ],
            wavelength,
            mapping,
        )

        all_samples.append(
            df
        )

        sheet_infos.append(
            info
        )

        print(
            f"  Header row : {info['header_row']}"
        )

        print(
            f"  Time cols  : {info['time_columns']}"
        )

        print(
            f"  Time range : "
            f"{info['time_min']:.6f} -> "
            f"{info['time_max']:.6f}"
        )

        print(
            f"  Samples    : {len(df)}"
        )

    sample_df = pd.concat(
        all_samples,
        ignore_index=True,
    )

    # ------------------------------------------------------------------
    # Basic sample consistency
    # ------------------------------------------------------------------

    consistency_df = sheet_consistency(
        sample_df,
        mapping,
    )

    print_section(
        "SAMPLE / MAPPING CONSISTENCY"
    )

    print(
        consistency_df.to_string(
            index=False
        )
    )

    # ------------------------------------------------------------------
    # Year scale
    # ------------------------------------------------------------------

    year_summary_df = year_summary(
        sample_df
    )

    print_section(
        "YEAR RAW SCALE SUMMARY"
    )

    print(
        year_summary_df.to_string(
            index=False,
            float_format=lambda x: f"{x:.6g}",
        )
    )

    # ------------------------------------------------------------------
    # Scale ratios
    # ------------------------------------------------------------------

    ratios_df = year_scale_ratios(
        sample_df
    )

    print_section(
        "YEAR SCALE RATIOS"
    )

    print(
        ratios_df.to_string(
            index=False,
            float_format=lambda x: f"{x:.8g}",
        )
    )

    # ------------------------------------------------------------------
    # Sample scale profile
    # ------------------------------------------------------------------

    profile_df = sample_scale_profile(
        sample_df
    )

    profile_df = (
        calculate_year_normalized_scale(
            profile_df
        )
    )

    print_section(
        "SAMPLE SCALE PROFILE"
    )

    display_columns = [
        "SampleId",
        "MappedGroup",
        "Year",
    ]

    for wavelength in [250, 308, 440]:

        column = (
            f"MedianAbs_{wavelength}"
        )

        if column in profile_df.columns:
            display_columns.append(
                column
            )

    print(
        profile_df[
            display_columns
        ]
        .sort_values(
            [
                "Year",
                "SampleId",
            ]
        )
        .to_string(
            index=False,
            float_format=lambda x: f"{x:.6g}",
        )
    )

    # ------------------------------------------------------------------
    # Year-specific CV-style scale ranges
    # ------------------------------------------------------------------

    print_section(
        "YEAR SCALE RANGE / LOG10"
    )

    for wavelength in [250, 308, 440]:

        column = (
            f"MedianAbs_{wavelength}"
        )

        if column not in profile_df.columns:
            continue

        temp = profile_df[
            [
                "SampleId",
                "Year",
                column,
            ]
        ].copy()

        temp["Log10MedianAbs"] = np.nan

        valid = temp[
            temp[column] > 0
        ].index

        temp.loc[
            valid,
            "Log10MedianAbs",
        ] = np.log10(
            temp.loc[
                valid,
                column,
            ]
        )

        summary = (
            temp
            .groupby("Year")[
                [
                    column,
                    "Log10MedianAbs",
                ]
            ]
            .agg(
                [
                    "min",
                    "median",
                    "max",
                ]
            )
        )

        print(
            f"\n{wavelength} nm"
        )

        print(
            summary.to_string(
                float_format=lambda x: f"{x:.6g}"
            )
        )

    # ------------------------------------------------------------------
    # Group × Year within raw scale
    # ------------------------------------------------------------------

    print_section(
        "GROUP × YEAR RAW SCALE"
    )

    group_year = (
        sample_df
        .groupby(
            [
                "Wavelength",
                "MappedGroup",
                "Year",
            ]
        )[
            [
                "MedianAbs",
                "MaxAbs",
                "AbsAUC",
            ]
        ]
        .median()
        .reset_index()
        .sort_values(
            [
                "Wavelength",
                "MappedGroup",
                "Year",
            ]
        )
    )

    print(
        group_year.to_string(
            index=False,
            float_format=lambda x: f"{x:.6g}",
        )
    )

    # ------------------------------------------------------------------
    # Formula audit
    # ------------------------------------------------------------------

    print_section(
        "FORMULA AUDIT"
    )

    print(
        structure_df[
            [
                "Sheet",
                "FormulaCellsInSignalRegion",
                "FormulaCachedValueCells",
                "FormulaValueDifferenceCells",
            ]
        ].to_string(
            index=False
        )
    )

    # ------------------------------------------------------------------
    # Format audit
    # ------------------------------------------------------------------

    format_df = format_consistency(
        wb_formula
    )

    print_section(
        "CELL FORMAT / STYLE AUDIT"
    )

    print(
        format_df.to_string(
            index=False
        )
    )

    # ------------------------------------------------------------------
    # Signal integrity
    # ------------------------------------------------------------------

    print_section(
        "SIGNAL INTEGRITY"
    )

    integrity = (
        sample_df
        .groupby("Wavelength")[
            [
                "SignalPoints",
                "BlankCells",
                "NumericCells",
                "TextCells",
                "ErrorCells",
                "FormulaCells",
            ]
        ]
        .agg(
            [
                "min",
                "median",
                "max",
            ]
        )
    )

    print(
        integrity.to_string()
    )

    # ------------------------------------------------------------------
    # Duplicate pairs
    # ------------------------------------------------------------------

    print_section(
        "KNOWN DUPLICATE PAIRS"
    )

    for first, second in DUPLICATE_PAIRS:

        print(
            f"\nPair {first} == {second}"
        )

        for wavelength in [250, 308, 440]:

            current = sample_df[
                sample_df[
                    "Wavelength"
                ] == wavelength
            ]

            a = current[
                current[
                    "SampleId"
                ] == first
            ]

            b = current[
                current[
                    "SampleId"
                ] == second
            ]

            if (
                a.empty
                or b.empty
            ):
                print(
                    f"  {wavelength} nm: missing"
                )
                continue

            a = a.iloc[0]
            b = b.iloc[0]

            ratio = np.nan

            if (
                abs(
                    float(
                        a["MedianAbs"]
                    )
                ) > 1e-15
            ):
                ratio = (
                    float(
                        b["MedianAbs"]
                    )
                    / float(
                        a["MedianAbs"]
                    )
                )

            print(
                f"  {wavelength} nm: "
                f"MedianAbs "
                f"{a['MedianAbs']:.6g} vs "
                f"{b['MedianAbs']:.6g} | "
                f"ratio={ratio:.6g}"
            )

    # ------------------------------------------------------------------
    # Representative samples
    # ------------------------------------------------------------------

    representatives = representative_samples(
        sample_df
    )

    print_section(
        "REPRESENTATIVE LOW / MEDIAN / HIGH SAMPLES"
    )

    print(
        representatives.to_string(
            index=False,
            float_format=lambda x: f"{x:.6g}",
        )
    )

    # ------------------------------------------------------------------
    # Save all outputs
    # ------------------------------------------------------------------

    outputs = {
        "workbook_structure":
            OUTPUT_DIR
            / "provenance_workbook_structure.csv",

        "sample_metrics":
            OUTPUT_DIR
            / "provenance_sample_metrics.csv",

        "sample_scale_profile":
            OUTPUT_DIR
            / "provenance_sample_scale_profile.csv",

        "year_summary":
            OUTPUT_DIR
            / "provenance_year_summary.csv",

        "year_scale_ratios":
            OUTPUT_DIR
            / "provenance_year_scale_ratios.csv",

        "group_year_scale":
            OUTPUT_DIR
            / "provenance_group_year_scale.csv",

        "consistency":
            OUTPUT_DIR
            / "provenance_sample_consistency.csv",

        "format_audit":
            OUTPUT_DIR
            / "provenance_format_audit.csv",

        "representatives":
            OUTPUT_DIR
            / "provenance_representative_samples.csv",
    }

    structure_df.to_csv(
        outputs["workbook_structure"],
        index=False,
        encoding="utf-8-sig",
    )

    sample_df.to_csv(
        outputs["sample_metrics"],
        index=False,
        encoding="utf-8-sig",
    )

    profile_df.to_csv(
        outputs["sample_scale_profile"],
        index=False,
        encoding="utf-8-sig",
    )

    year_summary_df.to_csv(
        outputs["year_summary"],
        index=False,
        encoding="utf-8-sig",
    )

    ratios_df.to_csv(
        outputs["year_scale_ratios"],
        index=False,
        encoding="utf-8-sig",
    )

    group_year.to_csv(
        outputs["group_year_scale"],
        index=False,
        encoding="utf-8-sig",
    )

    consistency_df.to_csv(
        outputs["consistency"],
        index=False,
        encoding="utf-8-sig",
    )

    format_df.to_csv(
        outputs["format_audit"],
        index=False,
        encoding="utf-8-sig",
    )

    representatives.to_csv(
        outputs["representatives"],
        index=False,
        encoding="utf-8-sig",
    )

    # ------------------------------------------------------------------
    # Close
    # ------------------------------------------------------------------

    wb_formula.close()
    wb_values.close()

    print_section(
        "OUTPUTS"
    )

    for name, path in outputs.items():

        print(
            f"{name:24s}: {path}"
        )


if __name__ == "__main__":
    main()