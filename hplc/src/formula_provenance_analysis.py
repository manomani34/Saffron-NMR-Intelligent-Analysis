from __future__ import annotations

from pathlib import Path

import numpy as np
import openpyxl


ROOT = Path(__file__).resolve().parents[1]

EXCEL_PATH = ROOT / "data" / "raw" / "hplc.xlsx"
OUTPUT_DIR = ROOT / "outputs"

OUTPUT_DIR.mkdir(parents=True, exist_ok=True)


SHEET = "safranal 308"

SAMPLE39_ROW = 41
SAMPLE40_ROW = 42

EXPECTED_FACTOR = 0.53


def main():

    print("=" * 78)
    print("HPLC FORMULA PROVENANCE ANALYSIS")
    print("=" * 78)

    wb_formula = openpyxl.load_workbook(
        EXCEL_PATH,
        data_only=False,
        read_only=False,
    )

    wb_values = openpyxl.load_workbook(
        EXCEL_PATH,
        data_only=True,
        read_only=False,
    )

    ws_f = wb_formula[SHEET]
    ws_v = wb_values[SHEET]

    print(f"\nSheet: {SHEET}")

    # ------------------------------------------------------------------
    # Sample IDs
    # ------------------------------------------------------------------

    print("\n")
    print("=" * 78)
    print("SAMPLE ROW IDENTIFICATION")
    print("=" * 78)

    print(
        "Row 41 SampleId:",
        ws_v.cell(SAMPLE39_ROW, 3).value,
    )

    print(
        "Row 42 SampleId:",
        ws_v.cell(SAMPLE40_ROW, 3).value,
    )

    print(
        "Row 41 Group:",
        ws_v.cell(SAMPLE39_ROW, 4).value,
    )

    print(
        "Row 42 Group:",
        ws_v.cell(SAMPLE40_ROW, 4).value,
    )

    # ------------------------------------------------------------------
    # Compare all cells
    # ------------------------------------------------------------------

    ratios = []
    absolute_errors = []
    formula_cells = 0
    valid_pairs = 0
    invalid_pairs = 0

    formula_examples = []

    # E:1804 approximately corresponds to 1800 signal points.
    for col in range(
        5,
        ws_f.max_column + 1,
    ):

        formula_cell = ws_f.cell(
            SAMPLE40_ROW,
            col,
        )

        sample39_cell = ws_v.cell(
            SAMPLE39_ROW,
            col,
        )

        sample40_cell = ws_v.cell(
            SAMPLE40_ROW,
            col,
        )

        if formula_cell.data_type == "f":

            formula_cells += 1

            if len(formula_examples) < 15:

                formula_examples.append(
                    {
                        "cell": formula_cell.coordinate,
                        "formula": formula_cell.value,
                        "cached_value": sample40_cell.value,
                    }
                )

        v39 = sample39_cell.value
        v40 = sample40_cell.value

        try:

            v39 = float(v39)
            v40 = float(v40)

        except (
            TypeError,
            ValueError,
        ):

            invalid_pairs += 1
            continue

        if not (
            np.isfinite(v39)
            and np.isfinite(v40)
        ):
            invalid_pairs += 1
            continue

        # Avoid division by zero.
        if abs(v39) < 1e-15:

            invalid_pairs += 1
            continue

        valid_pairs += 1

        ratio = v40 / v39

        ratios.append(ratio)

        expected = (
            v39
            * EXPECTED_FACTOR
        )

        absolute_errors.append(
            abs(
                v40 - expected
            )
        )

    ratios = np.asarray(
        ratios,
        dtype=float,
    )

    absolute_errors = np.asarray(
        absolute_errors,
        dtype=float,
    )

    # ------------------------------------------------------------------
    # Formula examples
    # ------------------------------------------------------------------

    print("\n")
    print("=" * 78)
    print("FORMULA EXAMPLES")
    print("=" * 78)

    for item in formula_examples:

        print(
            f"{item['cell']}: "
            f"{item['formula']} "
            f"-> cached={item['cached_value']}"
        )

    # ------------------------------------------------------------------
    # Ratio statistics
    # ------------------------------------------------------------------

    print("\n")
    print("=" * 78)
    print("SAMPLE 40 / SAMPLE 39 RATIO")
    print("=" * 78)

    if len(ratios) == 0:

        print(
            "No valid numeric pairs found."
        )

    else:

        print(
            f"Formula cells on row 42 : "
            f"{formula_cells}"
        )

        print(
            f"Valid numeric pairs      : "
            f"{valid_pairs}"
        )

        print(
            f"Invalid / skipped pairs  : "
            f"{invalid_pairs}"
        )

        print(
            f"Expected factor         : "
            f"{EXPECTED_FACTOR:.8f}"
        )

        print(
            f"Median ratio            : "
            f"{np.median(ratios):.12f}"
        )

        print(
            f"Mean ratio              : "
            f"{np.mean(ratios):.12f}"
        )

        print(
            f"Std ratio               : "
            f"{np.std(ratios):.12e}"
        )

        print(
            f"Min ratio               : "
            f"{np.min(ratios):.12f}"
        )

        print(
            f"Max ratio               : "
            f"{np.max(ratios):.12f}"
        )

        print(
            f"Mean abs error vs 0.53  : "
            f"{np.mean(absolute_errors):.12e}"
        )

        print(
            f"Max abs error vs 0.53   : "
            f"{np.max(absolute_errors):.12e}"
        )

        expected_ratio_error = np.abs(
            ratios
            - EXPECTED_FACTOR
        )

        print(
            f"Median |ratio-0.53|    : "
            f"{np.median(expected_ratio_error):.12e}"
        )

        print(
            f"Max |ratio-0.53|       : "
            f"{np.max(expected_ratio_error):.12e}"
        )

        exact_fraction = float(
            np.mean(
                expected_ratio_error
                < 1e-10
            )
        )

        near_fraction = float(
            np.mean(
                expected_ratio_error
                < 1e-6
            )
        )

        print(
            f"Exact fraction (<1e-10): "
            f"{exact_fraction:.6f}"
        )

        print(
            f"Near fraction  (<1e-6) : "
            f"{near_fraction:.6f}"
        )

    # ------------------------------------------------------------------
    # Row 34 / formula anomaly
    # ------------------------------------------------------------------

    print("\n")
    print("=" * 78)
    print("ANOMALOUS FORMULA CELL")
    print("=" * 78)

    anomalous = []

    for col in range(
        1,
        ws_f.max_column + 1,
    ):

        cell = ws_f.cell(
            34,
            col,
        )

        if cell.data_type == "f":

            cached = ws_v.cell(
                34,
                col,
            ).value

            anomalous.append(
                (
                    cell.coordinate,
                    cell.value,
                    cached,
                )
            )

    if anomalous:

        for coordinate, formula, cached in anomalous:

            print(
                f"{coordinate}: "
                f"{formula} "
                f"-> cached={cached}"
            )

    else:

        print(
            "No formula found on row 34."
        )

    # ------------------------------------------------------------------
    # Save
    # ------------------------------------------------------------------

    rows = []

    for idx, ratio in enumerate(
        ratios,
        start=1,
    ):

        rows.append(
            {
                "SignalIndex": idx,
                "Ratio_Sample40_over_Sample39": ratio,
                "ExpectedFactor": EXPECTED_FACTOR,
                "AbsoluteRatioError": abs(
                    ratio
                    - EXPECTED_FACTOR
                ),
            }
        )

    ratio_df_path = (
        OUTPUT_DIR
        / "formula_provenance_sample39_40_ratio.csv"
    )

    if rows:

        import pandas as pd

        pd.DataFrame(rows).to_csv(
            ratio_df_path,
            index=False,
            encoding="utf-8-sig",
        )

        print("\n")
        print(
            f"Ratio output: "
            f"{ratio_df_path}"
        )

    wb_formula.close()
    wb_values.close()

    print("\n")
    print("=" * 78)
    print("DONE")
    print("=" * 78)


if __name__ == "__main__":
    main()