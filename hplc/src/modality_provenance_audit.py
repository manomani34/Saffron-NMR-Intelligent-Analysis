from __future__ import annotations

from pathlib import Path
from itertools import combinations
import numpy as np
import pandas as pd
import openpyxl


ROOT = Path(__file__).resolve().parents[1]

EXCEL_PATH = ROOT / "data" / "raw" / "hplc.xlsx"
MAPPING_PATH = ROOT / "data" / "mapping" / "sample_mapping.csv"
OUTPUT_DIR = ROOT / "outputs"

OUTPUT_DIR.mkdir(parents=True, exist_ok=True)


SHEETS = {
    "picrocrocin 250": 250,
    "safranal 308": 308,
    "Crocin 440": 440,
}

MAX_TIME = 30.0

MIN_VALID_POINTS = 1000

EXACT_REL_TOL = 1e-12

SCALE_RESIDUAL_TOL = 1e-8

EPS = 1e-15


# =============================================================================
# HELPERS
# =============================================================================

def safe_float(value):
    try:
        value = float(value)

        if np.isfinite(value):
            return value

    except Exception:
        pass

    return None


def norm(value) -> str:
    if value is None:
        return ""

    return str(value).strip()


def find_header_row(ws):

    for row in range(
        1,
        min(ws.max_row, 30) + 1,
    ):

        values = [
            norm(
                ws.cell(
                    row,
                    col,
                ).value
            ).lower()
            for col in range(
                1,
                min(ws.max_column, 20) + 1,
            )
        ]

        if (
            "number" in values
            and "group" in values
        ):
            return row

    raise ValueError(
        f"Header row not found: {ws.title}"
    )


def detect_column(
    ws,
    header_row,
    candidates,
    fallback,
):

    for col in range(
        1,
        ws.max_column + 1,
    ):

        text = norm(
            ws.cell(
                header_row,
                col,
            ).value
        ).lower()

        if text in candidates:
            return col

    return fallback


def detect_time_columns(
    ws,
    header_row,
):

    columns = []
    times = []

    for col in range(
        1,
        ws.max_column + 1,
    ):

        value = ws.cell(
            header_row,
            col,
        ).value

        numeric = safe_float(
            value
        )

        if numeric is None:
            continue

        if (
            numeric > 0
            and numeric <= MAX_TIME
        ):

            columns.append(col)
            times.append(numeric)

    order = np.argsort(
        np.asarray(times)
    )

    columns = [
        columns[i]
        for i in order
    ]

    times = np.asarray(
        [
            times[i]
            for i in order
        ],
        dtype=float,
    )

    return columns, times


def load_mapping():

    mapping = pd.read_csv(
        MAPPING_PATH
    )

    normalized = {
        str(c)
        .strip()
        .lower()
        .replace("_", "")
        .replace(" ", ""): c
        for c in mapping.columns
    }

    sample_col = None
    group_col = None
    year_col = None

    for key, col in normalized.items():

        if key in {
            "sampleid",
            "sample",
            "number",
        }:
            sample_col = col

        if key in {
            "group",
            "class",
            "region",
        }:
            group_col = col

        if key in {
            "year",
            "harvestyear",
            "harvest",
        }:
            year_col = col

    if sample_col is None:
        raise ValueError(
            "Sample ID column not found."
        )

    if group_col is None:
        raise ValueError(
            "Group column not found."
        )

    result = pd.DataFrame()

    result["SampleId"] = pd.to_numeric(
        mapping[sample_col],
        errors="coerce",
    )

    result["Group"] = (
        mapping[group_col]
        .astype(str)
        .str.strip()
    )

    result["Year"] = (
        mapping[year_col]
        .astype(str)
        .str.strip()
        if year_col is not None
        else ""
    )

    result = result.dropna(
        subset=["SampleId"]
    )

    result["SampleId"] = (
        result["SampleId"]
        .astype(int)
    )

    return result


# =============================================================================
# LOAD ONE SHEET
# =============================================================================

def load_sheet(
    ws_formula,
    ws_values,
):

    header = find_header_row(
        ws_formula
    )

    number_col = detect_column(
        ws_formula,
        header,
        {
            "number",
            "sample number",
        },
        3,
    )

    group_col = detect_column(
        ws_formula,
        header,
        {"group"},
        4,
    )

    time_columns, times = (
        detect_time_columns(
            ws_formula,
            header,
        )
    )

    signals = {}
    formula_counts = {}
    formula_examples = {}

    workbook_groups = {}

    for row in range(
        header + 1,
        ws_formula.max_row + 1,
    ):

        sample_value = ws_values.cell(
            row,
            number_col,
        ).value

        sample_id = safe_float(
            sample_value
        )

        if sample_id is None:
            continue

        sample_id = int(
            sample_id
        )

        values = np.full(
            len(time_columns),
            np.nan,
            dtype=float,
        )

        formulas = []

        for i, col in enumerate(
            time_columns
        ):

            formula_cell = ws_formula.cell(
                row,
                col,
            )

            value_cell = ws_values.cell(
                row,
                col,
            )

            value = safe_float(
                value_cell.value
            )

            if value is not None:
                values[i] = value

            if formula_cell.data_type == "f":

                formulas.append(
                    (
                        formula_cell.coordinate,
                        formula_cell.value,
                    )
                )

        signals[sample_id] = values

        formula_counts[
            sample_id
        ] = len(formulas)

        formula_examples[
            sample_id
        ] = formulas[:10]

        workbook_groups[
            sample_id
        ] = norm(
            ws_values.cell(
                row,
                group_col,
            ).value
        )

    return {
        "signals": signals,
        "formula_counts": formula_counts,
        "formula_examples": formula_examples,
        "groups": workbook_groups,
        "times": times,
        "header": header,
    }


# =============================================================================
# COMPARE TWO SIGNALS
# =============================================================================

def compare(
    a,
    b,
):

    valid = (
        np.isfinite(a)
        & np.isfinite(b)
    )

    n = int(
        valid.sum()
    )

    if n < MIN_VALID_POINTS:

        return {
            "valid": n,
            "exact": False,
            "alpha": np.nan,
            "residual": np.nan,
            "correlation": np.nan,
        }

    x = a[valid]
    y = b[valid]

    scale = max(
        np.max(np.abs(x)),
        np.max(np.abs(y)),
        1.0,
    )

    max_diff = np.max(
        np.abs(x - y)
    )

    exact = (
        max_diff
        <= EXACT_REL_TOL * scale
    )

    denom = float(
        np.dot(x, x)
    )

    if abs(denom) < EPS:

        alpha = np.nan
        residual = np.nan

    else:

        alpha = float(
            np.dot(x, y)
            / denom
        )

        residual = float(
            np.linalg.norm(
                y - alpha * x
            )
            / max(
                np.linalg.norm(y),
                EPS,
            )
        )

    if (
        np.std(x) < EPS
        or np.std(y) < EPS
    ):

        correlation = np.nan

    else:

        correlation = float(
            np.corrcoef(
                x,
                y,
            )[0, 1]
        )

    return {
        "valid": n,
        "exact": exact,
        "alpha": alpha,
        "residual": residual,
        "correlation": correlation,
    }


# =============================================================================
# MAIN
# =============================================================================

def main():

    print("=" * 78)
    print(
        "HPLC MODALITY-LEVEL PROVENANCE AUDIT"
    )
    print("=" * 78)

    mapping = load_mapping()

    wb_formula = openpyxl.load_workbook(
        EXCEL_PATH,
        data_only=False,
        read_only=False,
    )

    wb_values = openpyxl.load_workbook(
        EXCEL_PATH,
        data_only=True,
        read_only=False,
    )

    all_relation_rows = []
    all_formula_rows = []
    all_integrity_rows = []

    # -------------------------------------------------------------------------
    # Each wavelength independently
    # -------------------------------------------------------------------------

    for sheet_name, wavelength in SHEETS.items():

        print("\n")
        print("-" * 78)
        print(
            f"{wavelength} nm | {sheet_name}"
        )
        print("-" * 78)

        data = load_sheet(
            wb_formula[sheet_name],
            wb_values[sheet_name],
        )

        signals = data[
            "signals"
        ]

        sample_ids = sorted(
            signals.keys()
        )

        print(
            f"Samples: {len(sample_ids)}"
        )

        print(
            f"Points : {len(data['times'])}"
        )

        # ---------------------------------------------------------------------
        # Integrity
        # ---------------------------------------------------------------------

        for sample_id in sample_ids:

            signal = signals[
                sample_id
            ]

            valid = np.isfinite(
                signal
            )

            all_integrity_rows.append(
                {
                    "Wavelength": wavelength,
                    "SampleId": sample_id,
                    "TotalPoints": len(signal),
                    "ValidPoints": int(
                        valid.sum()
                    ),
                    "InvalidPoints": int(
                        (~valid).sum()
                    ),
                    "FormulaCells": data[
                        "formula_counts"
                    ].get(
                        sample_id,
                        0,
                    ),
                }
            )

        # ---------------------------------------------------------------------
        # Formula rows
        # ---------------------------------------------------------------------

        for sample_id in sample_ids:

            formula_count = data[
                "formula_counts"
            ].get(
                sample_id,
                0,
            )

            if formula_count > 0:

                examples = data[
                    "formula_examples"
                ][sample_id]

                all_formula_rows.append(
                    {
                        "Wavelength": wavelength,
                        "SampleId": sample_id,
                        "FormulaCells": formula_count,
                        "FormulaExamples": str(
                            examples
                        ),
                    }
                )

        # ---------------------------------------------------------------------
        # Pairwise audit
        # ---------------------------------------------------------------------

        exact_count = 0
        scaled_count = 0

        for sample_a, sample_b in combinations(
            sample_ids,
            2,
        ):

            result = compare(
                signals[sample_a],
                signals[sample_b],
            )

            if result["exact"]:

                exact_count += 1

                all_relation_rows.append(
                    {
                        "Wavelength": wavelength,
                        "SampleA": sample_a,
                        "SampleB": sample_b,
                        "Relation": "ExactDuplicate",
                        "Alpha": 1.0,
                        "Residual": 0.0,
                        "Correlation": 1.0,
                    }
                )

                continue

            if (
                np.isfinite(
                    result["alpha"]
                )
                and np.isfinite(
                    result["residual"]
                )
                and result["alpha"] > 0
                and result["residual"]
                <= SCALE_RESIDUAL_TOL
            ):

                scaled_count += 1

                all_relation_rows.append(
                    {
                        "Wavelength": wavelength,
                        "SampleA": sample_a,
                        "SampleB": sample_b,
                        "Relation": "ScaledCopy",
                        "Alpha": result["alpha"],
                        "Residual": result["residual"],
                        "Correlation": result["correlation"],
                    }
                )

        print(
            f"Exact duplicate pairs: "
            f"{exact_count}"
        )

        print(
            f"Scaled-copy pairs: "
            f"{scaled_count}"
        )

    # -------------------------------------------------------------------------
    # Outputs
    # -------------------------------------------------------------------------

    relations_df = pd.DataFrame(
        all_relation_rows
    )

    formula_df = pd.DataFrame(
        all_formula_rows
    )

    integrity_df = pd.DataFrame(
        all_integrity_rows
    )

    # -------------------------------------------------------------------------
    # Print relations
    # -------------------------------------------------------------------------

    print("\n")
    print("=" * 78)
    print(
        "ALL MODALITY-LEVEL RELATIONS"
    )
    print("=" * 78)

    if relations_df.empty:

        print(
            "No exact/scaled relations found."
        )

    else:

        print(
            relations_df.to_string(
                index=False,
                float_format=lambda x: f"{x:.12g}",
            )
        )

    # -------------------------------------------------------------------------
    # Formula audit
    # -------------------------------------------------------------------------

    print("\n")
    print("=" * 78)
    print(
        "FORMULA-BEARING SAMPLES"
    )
    print("=" * 78)

    if formula_df.empty:

        print(
            "No formula-bearing signal rows."
        )

    else:

        print(
            formula_df.to_string(
                index=False
            )
        )

    # -------------------------------------------------------------------------
    # Invalid values
    # -------------------------------------------------------------------------

    print("\n")
    print("=" * 78)
    print(
        "INVALID SIGNAL CELLS"
    )
    print("=" * 78)

    invalid_df = integrity_df[
        integrity_df[
            "InvalidPoints"
        ] > 0
    ]

    if invalid_df.empty:

        print(
            "None."
        )

    else:

        print(
            invalid_df.to_string(
                index=False
            )
        )

    # -------------------------------------------------------------------------
    # Save
    # -------------------------------------------------------------------------

    paths = {
        "relations":
            OUTPUT_DIR
            / "modality_provenance_relations.csv",

        "formulas":
            OUTPUT_DIR
            / "modality_provenance_formula_samples.csv",

        "integrity":
            OUTPUT_DIR
            / "modality_provenance_integrity.csv",
    }

    relations_df.to_csv(
        paths["relations"],
        index=False,
        encoding="utf-8-sig",
    )

    formula_df.to_csv(
        paths["formulas"],
        index=False,
        encoding="utf-8-sig",
    )

    integrity_df.to_csv(
        paths["integrity"],
        index=False,
        encoding="utf-8-sig",
    )

    wb_formula.close()
    wb_values.close()

    print("\n")
    print("=" * 78)
    print("OUTPUTS")
    print("=" * 78)

    for name, path in paths.items():

        print(
            f"{name:12s}: {path}"
        )


if __name__ == "__main__":
    main()