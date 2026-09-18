from __future__ import annotations

from pathlib import Path
from itertools import combinations
import re

import numpy as np
import pandas as pd
import openpyxl


# ==============================================================================
# PATHS
# ==============================================================================

ROOT = Path(__file__).resolve().parents[1]

EXCEL_PATH = ROOT / "data" / "raw" / "hplc.xlsx"
MAPPING_PATH = ROOT / "data" / "mapping" / "sample_mapping.csv"
OUTPUT_DIR = ROOT / "outputs"

OUTPUT_DIR.mkdir(parents=True, exist_ok=True)


# ==============================================================================
# SETTINGS
# ==============================================================================

SHEETS = {
    "Crocin 440": 440,
    "picrocrocin 250": 250,
    "safranal 308": 308,
}

MAX_TIME = 30.0

# Minimum valid points required in each wavelength
MIN_VALID_POINTS = 1000

# Exact duplicate threshold
EXACT_REL_TOL = 1e-12

# Derived-scaled-copy thresholds
SCALE_RESIDUAL_TOL = 1e-8
SCALE_FACTOR_CV_TOL = 1e-8

# Avoid numerical problems
EPS = 1e-15


# ==============================================================================
# HELPERS
# ==============================================================================

def safe_float(value):
    try:
        value = float(value)

        if np.isfinite(value):
            return value

    except Exception:
        pass

    return None


def normalize_text(value) -> str:
    if value is None:
        return ""

    return str(value).strip()


def find_header_row(ws) -> int:

    for row_idx in range(
        1,
        min(ws.max_row, 30) + 1,
    ):

        values = [
            normalize_text(
                ws.cell(
                    row_idx,
                    col_idx,
                ).value
            ).lower()
            for col_idx in range(
                1,
                min(ws.max_column, 20) + 1,
            )
        ]

        has_number = any(
            x in {
                "number",
                "sample number",
            }
            for x in values
        )

        has_group = "group" in values

        if has_number and has_group:
            return row_idx

    raise ValueError(
        f"Header row not found in '{ws.title}'."
    )


def detect_column(
    ws,
    header_row: int,
    candidates: set[str],
    fallback: int,
) -> int:

    for col_idx in range(
        1,
        ws.max_column + 1,
    ):

        text = (
            normalize_text(
                ws.cell(
                    header_row,
                    col_idx,
                ).value
            )
            .lower()
        )

        if text in candidates:
            return col_idx

    return fallback


def detect_time_columns(
    ws,
    header_row: int,
) -> tuple[list[int], np.ndarray]:

    columns = []
    times = []

    for col_idx in range(
        1,
        ws.max_column + 1,
    ):

        value = ws.cell(
            header_row,
            col_idx,
        ).value

        numeric = safe_float(value)

        if numeric is None:
            continue

        if (
            numeric > 0
            and numeric <= MAX_TIME
        ):
            columns.append(col_idx)
            times.append(numeric)

    if not columns:
        raise ValueError(
            f"No time columns found in '{ws.title}'."
        )

    order = np.argsort(
        np.asarray(times)
    )

    columns = [
        columns[i]
        for i in order
    ]

    times = np.asarray(
        [
            times[i]
            for i in order
        ],
        dtype=float,
    )

    return columns, times


# ==============================================================================
# MAPPING
# ==============================================================================

def load_mapping() -> pd.DataFrame:

    mapping = pd.read_csv(
        MAPPING_PATH
    )

    normalized = {
        str(col)
        .strip()
        .lower()
        .replace("_", "")
        .replace(" ", ""): col
        for col in mapping.columns
    }

    sample_col = None
    group_col = None
    year_col = None

    for name, real_col in normalized.items():

        if name in {
            "sampleid",
            "sample",
            "number",
        }:
            sample_col = real_col

        if name in {
            "group",
            "class",
            "region",
        }:
            group_col = real_col

        if name in {
            "year",
            "harvestyear",
            "harvest",
        }:
            year_col = real_col

    if sample_col is None:
        raise ValueError(
            "Sample ID column not found in mapping."
        )

    if group_col is None:
        raise ValueError(
            "Group column not found in mapping."
        )

    result = pd.DataFrame()

    result["SampleId"] = pd.to_numeric(
        mapping[sample_col],
        errors="coerce",
    )

    result["ActualClass"] = (
        mapping[group_col]
        .astype(str)
        .str.strip()
    )

    if year_col is not None:
        result["Year"] = (
            mapping[year_col]
            .astype(str)
            .str.strip()
        )
    else:
        result["Year"] = ""

    result = result.dropna(
        subset=["SampleId"]
    )

    result["SampleId"] = (
        result["SampleId"]
        .astype(int)
    )

    return result


# ==============================================================================
# SHEET EXTRACTION
# ==============================================================================

def extract_sheet(
    ws_formula,
    ws_values,
    wavelength: int,
) -> dict:

    header_row = find_header_row(
        ws_formula
    )

    number_col = detect_column(
        ws_formula,
        header_row,
        {
            "number",
            "sample number",
        },
        fallback=3,
    )

    group_col = detect_column(
        ws_formula,
        header_row,
        {"group"},
        fallback=4,
    )

    time_columns, times = (
        detect_time_columns(
            ws_formula,
            header_row,
        )
    )

    records = {}

    formula_info = {}

    for row_idx in range(
        header_row + 1,
        ws_formula.max_row + 1,
    ):

        sample_value = (
            ws_values.cell(
                row_idx,
                number_col,
            ).value
        )

        sample_numeric = safe_float(
            sample_value
        )

        if sample_numeric is None:
            continue

        sample_id = int(
            sample_numeric
        )

        intensity = np.full(
            len(time_columns),
            np.nan,
            dtype=float,
        )

        formulas = []

        for i, col_idx in enumerate(
            time_columns
        ):

            formula_cell = (
                ws_formula.cell(
                    row_idx,
                    col_idx,
                )
            )

            value_cell = (
                ws_values.cell(
                    row_idx,
                    col_idx,
                )
            )

            value = value_cell.value

            numeric = safe_float(
                value
            )

            if numeric is not None:
                intensity[i] = numeric

            if (
                formula_cell.data_type
                == "f"
            ):
                formulas.append(
                    {
                        "cell": formula_cell.coordinate,
                        "formula": formula_cell.value,
                        "cached": value,
                    }
                )

        records[sample_id] = intensity

        formula_info[sample_id] = {
            "row": row_idx,
            "group": normalize_text(
                ws_values.cell(
                    row_idx,
                    group_col,
                ).value
            ),
            "formulas": formulas,
        }

    return {
        "wavelength": wavelength,
        "header_row": header_row,
        "times": times,
        "records": records,
        "formula_info": formula_info,
    }


# ==============================================================================
# LOAD ALL WAVES
# ==============================================================================

def load_workbook_data():

    wb_formula = openpyxl.load_workbook(
        EXCEL_PATH,
        data_only=False,
        read_only=False,
    )

    wb_values = openpyxl.load_workbook(
        EXCEL_PATH,
        data_only=True,
        read_only=False,
    )

    result = {}

    for sheet_name, wavelength in SHEETS.items():

        if sheet_name not in wb_formula.sheetnames:
            raise ValueError(
                f"Missing sheet: {sheet_name}"
            )

        result[wavelength] = extract_sheet(
            wb_formula[sheet_name],
            wb_values[sheet_name],
            wavelength,
        )

    return (
        result,
        wb_formula,
        wb_values,
    )


# ==============================================================================
# SAMPLE METADATA
# ==============================================================================

def build_sample_metadata(
    mapping: pd.DataFrame,
    sheets: dict,
) -> pd.DataFrame:

    sample_ids = set(
        mapping["SampleId"]
        .astype(int)
    )

    for info in sheets.values():
        sample_ids &= set(
            info["records"].keys()
        )

    rows = []

    mapping_index = (
        mapping
        .drop_duplicates(
            subset=["SampleId"]
        )
        .set_index("SampleId")
    )

    for sample_id in sorted(
        sample_ids
    ):

        rows.append(
            {
                "SampleId": sample_id,
                "MappedGroup": (
                    mapping_index.loc[
                        sample_id,
                        "ActualClass",
                    ]
                    if sample_id
                    in mapping_index.index
                    else ""
                ),
                "Year": (
                    mapping_index.loc[
                        sample_id,
                        "Year",
                    ]
                    if sample_id
                    in mapping_index.index
                    else ""
                ),
            }
        )

    return pd.DataFrame(rows)


# ==============================================================================
# CELL INTEGRITY
# ==============================================================================

def calculate_cell_integrity(
    sheets: dict,
) -> pd.DataFrame:

    rows = []

    for wavelength, info in sheets.items():

        for sample_id, signal in (
            info["records"].items()
        ):

            valid = np.isfinite(
                signal
            )

            rows.append(
                {
                    "SampleId": sample_id,
                    "Wavelength": wavelength,
                    "TotalSignalPoints": len(
                        signal
                    ),
                    "ValidNumericPoints": int(
                        valid.sum()
                    ),
                    "InvalidPoints": int(
                        (~valid).sum()
                    ),
                }
            )

    return pd.DataFrame(rows)


# ==============================================================================
# PAIR COMPARISON
# ==============================================================================

def compare_signals(
    a: np.ndarray,
    b: np.ndarray,
) -> dict:

    valid = (
        np.isfinite(a)
        & np.isfinite(b)
    )

    n = int(
        valid.sum()
    )

    if n < MIN_VALID_POINTS:
        return {
            "valid_points": n,
            "exact": False,
            "alpha": np.nan,
            "relative_residual": np.nan,
            "correlation": np.nan,
        }

    x = a[valid].astype(float)
    y = b[valid].astype(float)

    # --------------------------------------------------------------
    # Exact duplicate test
    # --------------------------------------------------------------

    scale = max(
        float(np.max(np.abs(x))),
        float(np.max(np.abs(y))),
        1.0,
    )

    max_abs_diff = float(
        np.max(
            np.abs(x - y)
        )
    )

    exact = (
        max_abs_diff
        <= EXACT_REL_TOL * scale
    )

    # --------------------------------------------------------------
    # Scalar relationship y = alpha * x
    # --------------------------------------------------------------

    denom = float(
        np.dot(x, x)
    )

    if abs(denom) < EPS:
        alpha = np.nan
        residual = np.nan
    else:

        alpha = float(
            np.dot(x, y)
            / denom
        )

        residual_vector = (
            y
            - alpha * x
        )

        residual = float(
            np.linalg.norm(
                residual_vector
            )
            / max(
                np.linalg.norm(y),
                EPS,
            )
        )

    # --------------------------------------------------------------
    # Correlation
    # --------------------------------------------------------------

    if (
        np.std(x) < EPS
        or np.std(y) < EPS
    ):
        correlation = np.nan
    else:
        correlation = float(
            np.corrcoef(
                x,
                y,
            )[0, 1]
        )

    return {
        "valid_points": n,
        "exact": exact,
        "alpha": alpha,
        "relative_residual": residual,
        "correlation": correlation,
    }


# ==============================================================================
# FORMULA SOURCE DETECTION
# ==============================================================================

def formula_provenance(
    formula_list: list[dict],
) -> dict:

    if not formula_list:
        return {
            "formula_count": 0,
            "reference_rows": [],
            "multipliers": [],
        }

    reference_rows = []
    multipliers = []

    # Examples:
    # =E41*0.53
    # =F41*0.53
    # =E41
    pattern = re.compile(
        r"=\$?[A-Z]{1,3}\$?(\d+)"
        r"(?:\*([+-]?\d+(?:\.\d+)?(?:[eE][+-]?\d+)?))?",
        re.IGNORECASE,
    )

    for item in formula_list:

        formula = str(
            item["formula"]
        ).strip()

        match = pattern.search(
            formula
        )

        if not match:
            continue

        try:
            reference_rows.append(
                int(
                    match.group(1)
                )
            )
        except Exception:
            pass

        if match.group(2) is not None:

            try:
                multipliers.append(
                    float(
                        match.group(2)
                    )
                )
            except Exception:
                pass

    return {
        "formula_count": len(
            formula_list
        ),
        "reference_rows": sorted(
            set(reference_rows)
        ),
        "multipliers": multipliers,
    }


# ==============================================================================
# PAIRWISE RELATION AUDIT
# ==============================================================================

def audit_pairwise_relations(
    sample_ids: list[int],
    sheets: dict,
    metadata: pd.DataFrame,
) -> tuple[
    pd.DataFrame,
    pd.DataFrame,
]:

    duplicate_rows = []
    relation_rows = []

    metadata_index = (
        metadata
        .set_index("SampleId")
    )

    print(
        "\nScanning sample pairs..."
    )

    pair_counter = 0

    total_pairs = (
        len(sample_ids)
        * (
            len(sample_ids) - 1
        )
        // 2
    )

    for sample_a, sample_b in combinations(
        sample_ids,
        2,
    ):

        pair_counter += 1

        wavelength_results = {}

        exact_all = True

        alphas = []
        residuals = []
        correlations = []
        valid_counts = []

        for wavelength, info in sheets.items():

            a = info[
                "records"
            ][sample_a]

            b = info[
                "records"
            ][sample_b]

            result = compare_signals(
                a,
                b,
            )

            wavelength_results[
                wavelength
            ] = result

            valid_counts.append(
                result[
                    "valid_points"
                ]
            )

            if not result["exact"]:
                exact_all = False

            if np.isfinite(
                result["alpha"]
            ):
                alphas.append(
                    result["alpha"]
                )

            if np.isfinite(
                result[
                    "relative_residual"
                ]
            ):
                residuals.append(
                    result[
                        "relative_residual"
                    ]
                )

            if np.isfinite(
                result["correlation"]
            ):
                correlations.append(
                    result["correlation"]
                )

        # --------------------------------------------------------------
        # Exact duplicate across all three wavelengths
        # --------------------------------------------------------------

        if (
            exact_all
            and len(wavelength_results)
            == len(SHEETS)
        ):

            duplicate_rows.append(
                {
                    "SampleA": sample_a,
                    "SampleB": sample_b,
                    "GroupA": metadata_index.loc[
                        sample_a,
                        "MappedGroup",
                    ],
                    "GroupB": metadata_index.loc[
                        sample_b,
                        "MappedGroup",
                    ],
                    "YearA": metadata_index.loc[
                        sample_a,
                        "Year",
                    ],
                    "YearB": metadata_index.loc[
                        sample_b,
                        "Year",
                    ],
                    "Valid250": wavelength_results[
                        250
                    ]["valid_points"],
                    "Valid308": wavelength_results[
                        308
                    ]["valid_points"],
                    "Valid440": wavelength_results[
                        440
                    ]["valid_points"],
                }
            )

        # --------------------------------------------------------------
        # Scaled copy candidate
        #
        # We require a scalar relation in ALL 3 wavelengths.
        # Direction:
        # SampleB = alpha * SampleA
        # --------------------------------------------------------------

        if (
            len(alphas) == 3
            and len(residuals) == 3
        ):

            alpha_array = np.asarray(
                alphas,
                dtype=float,
            )

            residual_array = np.asarray(
                residuals,
                dtype=float,
            )

            mean_alpha = float(
                np.mean(
                    alpha_array
                )
            )

            factor_cv = float(
                np.std(
                    alpha_array
                )
                / max(
                    abs(mean_alpha),
                    EPS,
                )
            )

            mean_residual = float(
                np.mean(
                    residual_array
                )
            )

            max_residual = float(
                np.max(
                    residual_array
                )
            )

            # Positive scalar copies are expected
            positive_factor = (
                mean_alpha > 0
            )

            scaled_candidate = (
                positive_factor
                and factor_cv
                <= SCALE_FACTOR_CV_TOL
                and max_residual
                <= SCALE_RESIDUAL_TOL
            )

            if scaled_candidate:

                relation_rows.append(
                    {
                        "SampleA": sample_a,
                        "SampleB": sample_b,
                        "GroupA": metadata_index.loc[
                            sample_a,
                            "MappedGroup",
                        ],
                        "GroupB": metadata_index.loc[
                            sample_b,
                            "MappedGroup",
                        ],
                        "YearA": metadata_index.loc[
                            sample_a,
                            "Year",
                        ],
                        "YearB": metadata_index.loc[
                            sample_b,
                            "Year",
                        ],

                        "Alpha250": wavelength_results[
                            250
                        ]["alpha"],

                        "Alpha308": wavelength_results[
                            308
                        ]["alpha"],

                        "Alpha440": wavelength_results[
                            440
                        ]["alpha"],

                        "MeanAlpha": mean_alpha,
                        "AlphaCV": factor_cv,

                        "Residual250": wavelength_results[
                            250
                        ]["relative_residual"],

                        "Residual308": wavelength_results[
                            308
                        ]["relative_residual"],

                        "Residual440": wavelength_results[
                            440
                        ]["relative_residual"],

                        "MaxResidual": max_residual,

                        "Correlation250": wavelength_results[
                            250
                        ]["correlation"],

                        "Correlation308": wavelength_results[
                            308
                        ]["correlation"],

                        "Correlation440": wavelength_results[
                            440
                        ]["correlation"],
                    }
                )

        if (
            pair_counter % 100 == 0
            or pair_counter == total_pairs
        ):

            print(
                f"  {pair_counter}/{total_pairs}"
            )

    return (
        pd.DataFrame(
            duplicate_rows
        ),
        pd.DataFrame(
            relation_rows
        ),
    )


# ==============================================================================
# SAMPLE STATUS
# ==============================================================================

def build_sample_status(
    metadata: pd.DataFrame,
    integrity_df: pd.DataFrame,
    duplicate_df: pd.DataFrame,
    scaled_df: pd.DataFrame,
    sheets: dict,
) -> pd.DataFrame:

    rows = []

    invalid_lookup = (
        integrity_df
        .groupby("SampleId")[
            "InvalidPoints"
        ]
        .sum()
        .to_dict()
    )

    duplicate_ids = set(
        duplicate_df[
            "SampleA"
        ].tolist()
        + duplicate_df[
            "SampleB"
        ].tolist()
    )

    scaled_source_ids = set()
    scaled_derived_ids = set()

    for _, row in scaled_df.iterrows():

        scaled_source_ids.add(
            int(row["SampleA"])
        )

        scaled_derived_ids.add(
            int(row["SampleB"])
        )

    for sample_id in sorted(
        metadata["SampleId"]
        .astype(int)
        .tolist()
    ):

        status = "Independent"

        if sample_id in duplicate_ids:
            status = "ExactDuplicate"

        elif sample_id in scaled_derived_ids:
            status = "DerivedScaled"

        elif sample_id in scaled_source_ids:
            status = "SourceOfDerivedScaled"

        invalid_points = int(
            invalid_lookup.get(
                sample_id,
                0,
            )
        )

        if invalid_points > 0:
            if status == "Independent":
                status = (
                    "IndependentWithInvalidCells"
                )
            else:
                status = (
                    status
                    + "+InvalidCells"
                )

        formula_count = 0
        formula_sources = []
        formula_multipliers = []

        for wavelength, info in sheets.items():

            formula_info = info[
                "formula_info"
            ].get(
                sample_id,
                {},
            )

            formulas = formula_info.get(
                "formulas",
                [],
            )

            provenance = formula_provenance(
                formulas
            )

            formula_count += provenance[
                "formula_count"
            ]

            formula_sources.extend(
                provenance[
                    "reference_rows"
                ]
            )

            formula_multipliers.extend(
                provenance[
                    "multipliers"
                ]
            )

        rows.append(
            {
                "SampleId": sample_id,
                "MappedGroup": metadata.loc[
                    metadata["SampleId"]
                    == sample_id,
                    "MappedGroup",
                ].iloc[0],
                "Year": metadata.loc[
                    metadata["SampleId"]
                    == sample_id,
                    "Year",
                ].iloc[0],
                "Status": status,
                "InvalidSignalCells": invalid_points,
                "FormulaCellsAcrossSheets": formula_count,
                "FormulaReferenceRows": (
                    ",".join(
                        map(
                            str,
                            sorted(
                                set(
                                    formula_sources
                                )
                            ),
                        )
                    )
                ),
                "FormulaMultipliers": (
                    ",".join(
                        f"{x:.12g}"
                        for x in sorted(
                            set(
                                formula_multipliers
                            )
                        )
                    )
                ),
            }
        )

    return pd.DataFrame(rows)


# ==============================================================================
# SUMMARY
# ==============================================================================

def build_summary(
    metadata: pd.DataFrame,
    integrity_df: pd.DataFrame,
    duplicate_df: pd.DataFrame,
    scaled_df: pd.DataFrame,
    status_df: pd.DataFrame,
) -> pd.DataFrame:

    rows = []

    status_counts = (
        status_df[
            "Status"
        ]
        .value_counts()
        .to_dict()
    )

    rows.append(
        {
            "RawSampleIds": int(
                len(metadata)
            ),
            "ExactDuplicatePairs": int(
                len(duplicate_df)
            ),
            "ScaledDerivedRelations": int(
                len(scaled_df)
            ),
            "ExactDuplicateSamples": int(
                status_counts.get(
                    "ExactDuplicate",
                    0,
                )
            ),
            "DerivedScaledSamples": int(
                sum(
                    1
                    for status in status_df[
                        "Status"
                    ]
                    if status.startswith(
                        "DerivedScaled"
                    )
                )
            ),
            "IndependentSamplesBeforeRelationCleanup": int(
                len(metadata)
            ),
            "IndependentSamplesAfterRelationCleanup": int(
                sum(
                    status.startswith(
                        "Independent"
                    )
                    for status in status_df[
                        "Status"
                    ]
                )
            ),
            "SamplesWithInvalidSignalCells": int(
                (
                    status_df[
                        "InvalidSignalCells"
                    ]
                    > 0
                ).sum()
            ),
            "TotalInvalidSignalCells": int(
                status_df[
                    "InvalidSignalCells"
                ].sum()
            ),
        }
    )

    return pd.DataFrame(rows)


# ==============================================================================
# MAIN
# ==============================================================================

def main():

    print("=" * 78)
    print(
        "HPLC SAMPLE INDEPENDENCE AUDIT"
    )
    print("=" * 78)

    # ------------------------------------------------------------------
    # Files
    # ------------------------------------------------------------------

    if not EXCEL_PATH.exists():
        raise FileNotFoundError(
            f"Workbook not found:\n{EXCEL_PATH}"
        )

    if not MAPPING_PATH.exists():
        raise FileNotFoundError(
            f"Mapping not found:\n{MAPPING_PATH}"
        )

    # ------------------------------------------------------------------
    # Load
    # ------------------------------------------------------------------

    mapping = load_mapping()

    print(
        f"\nMapping rows: "
        f"{len(mapping)}"
    )

    (
        sheets,
        wb_formula,
        wb_values,
    ) = load_workbook_data()

    # ------------------------------------------------------------------
    # Sheet overview
    # ------------------------------------------------------------------

    print("\n")
    print("=" * 78)
    print("SHEET OVERVIEW")
    print("=" * 78)

    for wavelength, info in (
        sorted(
            sheets.items()
        )
    ):

        print(
            f"{wavelength} nm"
        )

        print(
            f"  Time points : "
            f"{len(info['times'])}"
        )

        print(
            f"  Time range  : "
            f"{info['times'].min():.6f} -> "
            f"{info['times'].max():.6f}"
        )

        print(
            f"  Samples     : "
            f"{len(info['records'])}"
        )

        formula_count = sum(
            len(
                x["formulas"]
            )
            for x in info[
                "formula_info"
            ].values()
        )

        print(
            f"  Formula cells: "
            f"{formula_count}"
        )

    # ------------------------------------------------------------------
    # Common samples
    # ------------------------------------------------------------------

    metadata = build_sample_metadata(
        mapping,
        sheets,
    )

    sample_ids = sorted(
        metadata["SampleId"]
        .astype(int)
        .tolist()
    )

    print(
        f"\nCommon samples across "
        f"all 3 wavelengths: "
        f"{len(sample_ids)}"
    )

    # ------------------------------------------------------------------
    # Integrity
    # ------------------------------------------------------------------

    integrity_df = (
        calculate_cell_integrity(
            sheets
        )
    )

    integrity_summary = (
        integrity_df
        .groupby("Wavelength")[
            [
                "TotalSignalPoints",
                "ValidNumericPoints",
                "InvalidPoints",
            ]
        ]
        .agg(
            [
                "min",
                "median",
                "max",
            ]
        )
    )

    print("\n")
    print("=" * 78)
    print("SIGNAL INTEGRITY")
    print("=" * 78)

    print(
        integrity_summary.to_string()
    )

    invalid_df = (
        integrity_df[
            integrity_df[
                "InvalidPoints"
            ] > 0
        ]
        .merge(
            metadata,
            on="SampleId",
            how="left",
        )
    )

    print("\nInvalid sample cells:")

    if invalid_df.empty:
        print(
            "  None"
        )
    else:
        print(
            invalid_df.to_string(
                index=False
            )
        )

    # ------------------------------------------------------------------
    # Pairwise relations
    # ------------------------------------------------------------------

    (
        duplicate_df,
        scaled_df,
    ) = audit_pairwise_relations(
        sample_ids,
        sheets,
        metadata,
    )

    # ------------------------------------------------------------------
    # Print duplicates
    # ------------------------------------------------------------------

    print("\n")
    print("=" * 78)
    print("EXACT DUPLICATE PAIRS")
    print("=" * 78)

    if duplicate_df.empty:
        print(
            "No exact duplicate pairs found."
        )
    else:
        print(
            duplicate_df.to_string(
                index=False
            )
        )

    # ------------------------------------------------------------------
    # Print scaled copies
    # ------------------------------------------------------------------

    print("\n")
    print("=" * 78)
    print(
        "DERIVED / SCALED COPY RELATIONS"
    )
    print("=" * 78)

    if scaled_df.empty:
        print(
            "No high-confidence scaled-copy "
            "relations found."
        )
    else:

        print(
            scaled_df.to_string(
                index=False,
                float_format=lambda x: f"{x:.12g}",
            )
        )

    # ------------------------------------------------------------------
    # Sample status
    # ------------------------------------------------------------------

    status_df = build_sample_status(
        metadata,
        integrity_df,
        duplicate_df,
        scaled_df,
        sheets,
    )

    print("\n")
    print("=" * 78)
    print(
        "SAMPLE STATUS"
    )
    print("=" * 78)

    print(
        status_df.to_string(
            index=False
        )
    )

    # ------------------------------------------------------------------
    # Summary
    # ------------------------------------------------------------------

    summary_df = build_summary(
        metadata,
        integrity_df,
        duplicate_df,
        scaled_df,
        status_df,
    )

    print("\n")
    print("=" * 78)
    print(
        "INDEPENDENCE SUMMARY"
    )
    print("=" * 78)

    print(
        summary_df.to_string(
            index=False
        )
    )

    # ------------------------------------------------------------------
    # Group counts after cleanup
    # ------------------------------------------------------------------

    cleaned_status = status_df[
        status_df[
            "Status"
        ].str.startswith(
            "Independent"
        )
    ].copy()

    group_counts = (
        cleaned_status
        .groupby(
            "MappedGroup"
        )
        .size()
        .reset_index(
            name="IndependentSamples"
        )
        .sort_values(
            "MappedGroup"
        )
    )

    print("\n")
    print("=" * 78)
    print(
        "GROUP COUNTS AFTER RELATION CLEANUP"
    )
    print("=" * 78)

    print(
        group_counts.to_string(
            index=False
        )
    )

    # ------------------------------------------------------------------
    # Save
    # ------------------------------------------------------------------

    outputs = {
        "sample_integrity":
            OUTPUT_DIR
            / "sample_independence_signal_integrity.csv",

        "invalid_cells":
            OUTPUT_DIR
            / "sample_independence_invalid_cells.csv",

        "duplicate_pairs":
            OUTPUT_DIR
            / "sample_independence_duplicate_pairs.csv",

        "scaled_relations":
            OUTPUT_DIR
            / "sample_independence_scaled_relations.csv",

        "sample_status":
            OUTPUT_DIR
            / "sample_independence_status.csv",

        "summary":
            OUTPUT_DIR
            / "sample_independence_summary.csv",

        "group_counts":
            OUTPUT_DIR
            / "sample_independence_group_counts.csv",
    }

    integrity_df.to_csv(
        outputs["sample_integrity"],
        index=False,
        encoding="utf-8-sig",
    )

    invalid_df.to_csv(
        outputs["invalid_cells"],
        index=False,
        encoding="utf-8-sig",
    )

    duplicate_df.to_csv(
        outputs["duplicate_pairs"],
        index=False,
        encoding="utf-8-sig",
    )

    scaled_df.to_csv(
        outputs["scaled_relations"],
        index=False,
        encoding="utf-8-sig",
    )

    status_df.to_csv(
        outputs["sample_status"],
        index=False,
        encoding="utf-8-sig",
    )

    summary_df.to_csv(
        outputs["summary"],
        index=False,
        encoding="utf-8-sig",
    )

    group_counts.to_csv(
        outputs["group_counts"],
        index=False,
        encoding="utf-8-sig",
    )

    # ------------------------------------------------------------------
    # Close workbooks
    # ------------------------------------------------------------------

    wb_formula.close()
    wb_values.close()

    # ------------------------------------------------------------------
    # Output paths
    # ------------------------------------------------------------------

    print("\n")
    print("=" * 78)
    print("OUTPUTS")
    print("=" * 78)

    for name, path in outputs.items():

        print(
            f"{name:24s}: {path}"
        )


if __name__ == "__main__":
    main()